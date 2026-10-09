# ISA · SCORE Slides — (c) 2026 Francisco Javier Cavero López y Alejandro García Fernández
# (SCORE Lab, I3US, Universidad de Sevilla; Grupo ISA). Licencia CC BY 4.0 — ver LICENSE.
# Las presentaciones generadas con este kit pueden usarse libremente sin atribución (LICENSE, parte 2).
"""Builds the component library (one .pptx per reusable component) from the group's source decks.
Each library file keeps the original slides verbatim: shapes, media, Morph transitions and animation
timelines. Run from a folder that holds the source decks:  python build_library.py SRC_DIR OUT_DIR"""
import json
import os
import sys

from pptx import Presentation
from lxml import etree

from isa_deck import import_slides, _resolve_template, DEFAULT_TEMPLATE, NS

_args = [a for a in sys.argv[1:] if not a.startswith("--")]
SRC = _args[0] if len(_args) > 0 else "."
OUT = _args[1] if len(_args) > 1 else "library"
D = {
    "c26": "CAiSE_2026_-_FJCL.pptx", "c25": "CAISE25-AGF.pptx", "icwe": "ICWE25-AGF.pptx", "jcis": "JCIS-AGF-v2.pptx",
    "phd": "PhD_Forum_FEBRERO_2026_-_AGF.pptx", "sc25": "SCORE_2025_-_FJCL.pptx", "icsoc": "ICSOC24-SOC4AI-FJCL.pptx",
}

# id, category, name, deck, slides, description
CATALOG = [
    # --- Openers / hooks
    ("hook-pricing-question", "Openers", "Pricing hook question", "c26", [3, 4, 5],
     "Collage of real pricing pages → decision-maker ‘OPTIMIZE’ → value-cost with the guiding question → two real pricings (GitHub vs CircleCI) with highlighted rows. Morph between the three."),
    ("quote-hook", "Openers", "Authority quote", "phd", [3],
     "Centred quote with the key fragment in alert red and the author, used as the very first content slide."),
    ("statement-hooks", "Openers", "Statement hooks", "phd", [4, 5, 6],
     "One-word statement (‘Customization’), a claim with a giant crimson word (THOUSANDS) and the thesis title with ‘Pricing’ highlighted."),
    ("character-cases", "Openers", "Presenter character with cases", "c25", [25, 26, 27],
     "Illustrated presenter character at the left; mini pricing tables with red frames marking inconsistent cases, ‘Depends on / Excludes’ arrows. Character reused across slides."),
    ("thinking-character-story", "Openers", "Thinking character story", "sc25", [3, 4, 5, 6, 7, 8, 9, 10, 12],
     "Character with a thought cloud whose contents change step by step (travel booking → real SaaS brands → task graph). Zoom/Wipe entrances and Morph to the composition graph."),
    # --- Agenda
    ("agenda-arc", "Agenda", "Agenda arc (4 states)", "c25", [3, 23, 43, 50],
     "The real agenda with icons (start flag, bulb, checklist, checkered flag) on the crimson arc; one copy per section with the current circle filled."),
    ("toc-numbered", "Agenda", "Numbered ‘What is this about?’", "jcis", [3, 18],
     "Numbered table of contents with large crimson numerals; the second copy marks the current section."),
    ("toc-list", "Agenda", "Table of Contents (PhD Forum)", "phd", [2, 25],
     "Plain numbered TOC with nested items; current section in bold. Used before each section of a long talk."),
    # --- Pricing domain
    ("pricing-table-build", "Pricing", "Pricing table build → configuration space", "c26", [6, 7, 8, 9, 10, 11, 12],
     "Icons of features scattered → BASIC/PRO/BUSINESS table with prices → usage-limits legend (Relaxed/Moderate/Strict) → add-on card ($10 / month) → explosion into monitors with every subscription Ⓢ(…) → ‘Configuration Space’. All Morph."),
    ("pricing-build-user", "Pricing", "Pricing build with user and subscriptions", "icsoc", [2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12],
     "User at a monitor → empty plan table → features → limits → add-ons → fan of user subscriptions on monitors → users coloured per subscription (4:3 deck, centred)."),
    ("curse-of-customization", "Pricing", "Curse of customization fan", "c26", [13, 14],
     "Small pricing table whose plans fan out with crimson lines into the full list of subscriptions Ⓢ(…) once add-ons are added."),
    ("pricification-timeline", "Pricing", "Pricification process timeline", "phd", [31, 32, 33, 34],
     "Horizontal crimson timeline from a SaaS screen to a pricing table: milestones (identification of features, essential/key features, generation of the pricing structure) and a BPMN detail."),
    ("evaluation-ground-truth", "Pricing", "Ground truth vs extraction (TP/FP/FN/TN)", "icsoc", [28, 29, 30, 31, 32, 33],
     "Point-based evaluation hub (TP, FP, FN, TN) then ground-truth vs extracted pricing tables with crimson frames per case (4:3 deck, centred)."),
    # --- DevOps / processes
    ("devops-to-pricing-driven", "Processes", "DevOps loop → Pricing-driven DevOps", "c25", [13, 14, 15, 16, 17, 18, 19, 20, 21],
     "Traditional DevOps infinity loop with a provider-events timeline (pe1…pen) → effort chart → definition card over a grey veil → codebase with feature toggles → Pricing-driven Dev/Ops loop with its five research lines."),
    ("devops-numbered-overview", "Processes", "Pricing-driven DevOps research lines", "jcis", [48, 49, 50, 51, 52, 53, 54],
     "Numbered list of the five lines next to the loop; each copy highlights one line and swaps the illustration (metamodel, design, self-adaptation, testing, analysis)."),
    ("feature-toggles-codebase", "Processes", "Feature toggles in the codebase", "jcis", [14, 15, 16, 17],
     "Codebase files with toggles, toggle context database and pricing → evaluation space size ‘8’ → architecture with micro-services."),
    ("feature-toggle-anatomy", "Processes", "Anatomy of a feature toggle", "icwe", [22, 23, 24],
     "‘Feature Toggles’ statement → SaaS code file with toggles → zoom into the code with Context / Rule / Feature callouts."),
    ("evaluation-expression", "Processes", "Anatomy of an evaluation expression", "phd", [17],
     "Pricing table + code snippet with dashed frames; the expression is resolved against a plan (GOLD → PLATINUM) with crimson arrows."),
    ("three-step-flow", "Processes", "Three-card problem flow", "c26", [20],
     "Three rounded pastel cards (rose → white → blue) joined by thick crimson/teal arrows: Decision-makers → Semantic gap → PI tools."),
    ("pipeline-flow", "Processes", "Data-collection pipeline", "jcis", [20, 21],
     "GitHub → paper → SaaS cloud → YAML with crimson elbow arrows and counters (16, 3, 30 SaaS), then zoom to the charts."),
    # --- Architecture
    ("layers-zoom", "Architecture", "Layer stack → zoom into one layer", "c26", [21, 22, 23],
     "PIIP layers (accessibility / orchestration / capability) → capability layer zoomed with MCP server and tools → the layers beside the system logo."),
    ("architecture-spotlight", "Architecture", "Architecture walkthrough with spotlight", "c26", [24, 25, 26, 27, 28, 29, 30, 31],
     "Three-layer architecture image with a crimson frame that Morph moves from the user query to reasoning, tool invocation, MCP server, tools, synthesis and response; final copies isolate layers."),
    ("diagram-zoom-pan", "Architecture", "Diagram zoom and pan", "icsoc", [19, 20, 21, 22, 23],
     "Full architecture diagram (extractor / process engine / results modeler) → Morph zooms and pans across its stages (4:3 deck, centred)."),
    ("acronym-expansion", "Architecture", "Acronym expansion", "c25", [37, 38, 39],
     "‘How can we do this?’ → giant red acronym C S O P → Morph spreads it and lines connect each letter to its word."),
    ("proposal-stack", "Architecture", "Operations stack (filter / valid / model)", "c25", [40, 41, 42],
     "Proposal diagram (iPricing → data/model → CSOP) with crimson frames, then the layered operations stack with a solver logo."),
    # --- Evaluation & results
    ("rq-cards", "Evaluation", "Research-question cards", "c26", [32],
     "Three white RQ cards with crimson pill, label, question and grey note."),
    ("rq-circles", "Evaluation", "Research questions with numbered circles", "jcis", [19],
     "Numbered crimson circles with each research question beside them."),
    ("methodology-funnel", "Evaluation", "Methodology funnel", "c26", [34],
     "Participants and LLM logos → pink trapezoids (raw templates → manual review → structural filtering) → final set, with angled connectors."),
    ("coverage-panels", "Evaluation", "Coverage + interpretation panels", "c26", [35],
     "Big pink coverage card (54 %) with formula pill, interpretation card and ‘open gaps’ chips."),
    ("results-table", "Evaluation", "Results table with RQ answers", "c26", [37],
     "Table with crimson header row, then two pink answer cards for RQ2 / RQ3."),
    ("evaluation-stages", "Evaluation", "Two-stage evaluation rings", "c25", [44],
     "Green and blue ring diagrams for Stage 01 / Stage 02 with short descriptions."),
    ("stat-split-panels", "Evaluation", "Dataset / result split panels", "c25", [45, 46],
     "Slide split in pale yellow (dataset: stat bubbles + SaaS logos) and pink (result: +20 % with breakdown, then trophy 100 %)."),
    ("split-outcome", "Evaluation", "Split outcome arrows", "c25", [48, 49],
     "Pricing files → solver → crimson fork to 65 % / 35 % with pictograms; then the inconsistency cases."),
    ("pyramid-levels", "Evaluation", "Capability pyramid L1–L3", "jcis", [28, 29, 30, 31, 32],
     "Empty triangle → L1 → L2 → L3 crimson bands, each with a code callout (.js) appearing beside it."),
    ("big-stat-out-of", "Evaluation", "‘4 out of 18’ big stats", "icwe", [39],
     "Yellow/pink split with giant fractions (4 out of 18 tools) and tool logos."),
    ("tools-arcs", "Evaluation", "Tools on arcs", "icwe", [40],
     "Tool logos inside partial coloured rings with short verdicts (Togglz, Unleash, LaunchDarkly…)."),
    ("challenges-hub", "Evaluation", "Challenges hub A–F", "icsoc", [43],
     "Two central nodes (internal / external) branching to six lettered challenges (4:3 deck, centred)."),
    # --- Closing patterns
    ("rotating-wheel", "Closing", "Rotating future-work wheel", "c26", [40, 41, 42, 43, 44, 45, 47],
     "Pentagon hub (‘future work’) → coloured segments appear one by one and the wheel rotates with Morph to bring each research line to the top, with its guiding question; ends beside the system logo."),
    ("segment-wheel-build", "Closing", "Segment wheel build (framework)", "icwe", [26, 27, 28, 29, 30, 31],
     "Framework logo in a pentagon; segments grow around it one per step until the classification wheel is complete."),
    ("chevron-steps", "Closing", "Chevron numbered steps", "phd", [53, 54, 55, 56],
     "Conclusions over the loop, then crimson/amber chevrons 1-2-3 stacking future-work items with partner logos."),
    ("numbered-points-0104", "Closing", "Future work 01–04", "sc25", [42],
     "Big thin numerals 01–04 with crimson headings."),
    ("conclusions-icons", "Closing", "Conclusions with icons", "sc25", [43],
     "Three crimson icon tiles with short captions."),
    ("conclusions-recap", "Closing", "Conclusions recap", "c25", [51, 52, 53, 55],
     "Numbered contributions growing on the left while the right side re-shows the proposal with a crimson frame on the part being recalled."),
    ("ribbon-numbered", "Closing", "Numbered ribbons with QR", "jcis", [37],
     "Crimson ribbons 01/02/03 with white numbered discs and a QR code."),
    ("puzzle-pieces", "Closing", "Four-piece puzzle", "sc25", [41],
     "Crimson puzzle with four labelled callouts (icon bars) around it."),
    ("strike-replace", "Openers", "Strike and replace a term", "sc25", [30, 31, 32],
     "QACO / QoS-aware composition → struck through in crimson → PACO / Pricing-aware composition appears below."),
]


def strip_slides(prs):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


def summarize(slides):
    tr, fx = [], {}
    for s in slides:
        x = etree.tostring(s._element, encoding="unicode")
        if "p159:morph" in x:
            tr.append("morph")
        elif "<p:transition" in x and "<p:fade" in x:
            tr.append("fade")
        else:
            tr.append("cut")
        import re
        for pid, cls in re.findall(r'presetID="(\d+)" presetClass="(\w+)"', x):
            k = f"{cls}:{pid}"
            fx[k] = fx.get(k, 0) + 1
    return tr, fx


PRESET_NAMES = {"1": "appear", "2": "fly_in", "3": "blinds", "5": "checkerboard", "6": "circle", "9": "dissolve", "10": "fade",
                "12": "peek_in", "16": "split", "22": "wipe", "37": "rise_up", "42": "float_in", "53": "zoom"}

os.makedirs(OUT, exist_ok=True)
catalog = []
for cid, cat, name, deck, nums, desc in CATALOG:
    base = Presentation(_resolve_template(DEFAULT_TEMPLATE))
    strip_slides(base)
    src = os.path.join(SRC, D[deck])
    off = 1524000 if deck == "icsoc" else 0
    slides = import_slides(base, src, set(nums), x_offset=off, compress=False)  # lossless: media copied byte for byte
    if off:  # 4:3 slide centred on 16:9: white side masks keep the original 4:3 framing (zooms, bleeds)
        from pptx.util import Emu
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        for sl in slides:
            for nm, x in (("!!4x3 mask left", 0), ("!!4x3 mask right", off + 9144000)):
                m = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(0), Emu(off), Emu(6858000))
                m.fill.solid(); m.fill.fore_color.rgb = RGBColor(255, 255, 255)
                m.line.fill.background(); m.shadow.inherit = False; m.name = nm
    path = os.path.join(OUT, cid + ".pptx")
    base.save(path)
    # Lossless by default (the plugin keeps the original media). Pass --optimize to recompress media
    # to 2x their largest on-slide size, as the design-system store needs (see optimize_pptx.py).
    if "--optimize" in sys.argv:
        from optimize_pptx import optimize
        optimize(path)
    tr, fx = summarize(slides)
    catalog.append({
        "id": cid, "category": cat, "name": name, "description": desc, "slides": len(slides),
        "source": f"{D[deck]} slides {nums}", "aspect": "4:3 centred on 16:9" if off else "16:9",
        "transitions": tr,
        "effects": {f"{k.split(':')[0]}:{PRESET_NAMES.get(k.split(':')[1], k.split(':')[1])}": v for k, v in sorted(fx.items())},
        "file": f"kit/library/{cid}.pptx.b64.txt", "bytes": os.path.getsize(path),
    })
    print(f"{cid:28s} {len(slides):2d} slides {os.path.getsize(path)/1e6:5.2f} MB  {''.join(t[0] for t in tr)}")
json.dump(catalog, open(os.path.join(OUT, "catalog.json"), "w"), ensure_ascii=False, indent=1)
print("total MB", sum(c["bytes"] for c in catalog) / 1e6)
