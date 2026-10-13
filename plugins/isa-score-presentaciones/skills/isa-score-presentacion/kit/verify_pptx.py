# ISA · SCORE Slides — (c) 2026 Francisco Javier Cavero López y Alejandro García Fernández
# (SCORE Lab, I3US, Universidad de Sevilla; Grupo ISA). Licencia CC BY 4.0 — ver LICENSE.
"""Comprueba un .pptx antes de entregarlo y, opcionalmente, hace una hoja de miniaturas.

    python3 verify_pptx.py talk.pptx [--sheet sheet.png] [--expect N]

Por diapositiva: transición (Morph/Fade/—), nº de animaciones de entrada y salida, formas `!!` y,
en cada diapositiva con Morph, cuántos nombres de forma comparte con la anterior (Morph empareja por nombre; `!!` lo fuerza). Avisa de:
  - diapositivas de contenido sin ninguna animación ni Morph de entrada (estáticas);
  - Morph sin ningún nombre de forma en común con la anterior (no emparejará nada);
    - paginación con tamaño, posición, color o estilo distinto; una caja fuera del lienzo se admite solo durante un Morph;
  - nº de diapositivas distinto de --expect.
--sheet renderiza con LibreOffice (soffice + pdftoppm) y compone una hoja numerada (LibreOffice no
muestra Morph, pero sí composición, textos e imágenes).
"""
import os, re, subprocess, sys, tempfile
from pptx import Presentation
from lxml import etree


def _number_signature(shape):
    """Return the visual signature that must remain stable for slide numbers."""
    run = next((r for p in shape.text_frame.paragraphs for r in p.runs), None)
    if run is None:
        return None
    color = run.font.color
    rgb = None
    if color.type is not None:
        try:
            rgb = str(color.rgb)
        except (ValueError, AttributeError):
            rgb = str(color.type)
    return (shape.left, shape.top, shape.width, shape.height,
            run.font.size, run.font.name, run.font.bold, run.font.italic, rgb)


def _number_shape(slide):
    for shape in slide.shapes:
        if shape.name == "Slide number" and shape.has_text_frame:
            return shape
    return None


def _check_slide_numbers(prs):
    """Check visible number boxes and allow only Morph zoom boxes off-canvas."""
    baseline = None
    issues = []
    off_canvas = []
    for index, slide in enumerate(prs.slides, 1):
        shape = _number_shape(slide)
        if shape is None:
            continue
        outside = (shape.left < 0 or shape.top < 0 or
                   shape.left + shape.width > prs.slide_width or
                   shape.top + shape.height > prs.slide_height)
        xml = etree.tostring(slide._element, encoding="unicode")
        morph = "p159:morph" in xml
        if outside:
            if morph:
                off_canvas.append(index)
            else:
                issues.append((index, "fuera del lienzo sin Morph de ampliación"))
            continue
        signature = _number_signature(shape)
        if baseline is None:
            baseline = signature
        elif signature != baseline:
            issues.append((index, "tamaño, posición, color o estilo distinto"))
    return baseline, issues, off_canvas


def analyse(path):
    prs = Presentation(path)
    rows, prev = [], set()
    for i, s in enumerate(prs.slides, 1):
        x = etree.tostring(s._element, encoding="unicode")
        tr = "Morph" if "p159:morph" in x else ("Fade" if re.search(r"<p:transition[^>]*>\s*<p:fade", x) else "—")
        ent = len(re.findall(r'presetClass="entr"', x))
        ext = len(re.findall(r'presetClass="exit"', x))
        allnames = {sh.name for sh in s.shapes}
        names = {n for n in allnames if n.startswith("!!")}
        shared = len(allnames & prev) if tr == "Morph" else None
        hidden = s._element.get("show") == "0"
        rows.append((i, tr, ent, ext, len(names), shared, len(s.shapes), hidden))
        prev = allnames
    return rows


def sheet(path, out):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    subprocess.run(["pdftoppm", "-r", "20", "-png", pdf, os.path.join(tmp, "p")], check=True)
    from PIL import Image, ImageDraw
    fs = sorted((f for f in os.listdir(tmp) if f.startswith("p-")), key=lambda f: int(re.findall(r"(\d+)", f)[0]))
    ims = [Image.open(os.path.join(tmp, f)).convert("RGB") for f in fs]
    w, h, cols = ims[0].width, ims[0].height, 8
    rows = (len(ims) + cols - 1) // cols
    sh = Image.new("RGB", (cols * (w + 6) + 6, rows * (h + 6) + 6), (90, 90, 90))
    d = ImageDraw.Draw(sh)
    for k, im in enumerate(ims):
        x, y = 6 + (k % cols) * (w + 6), 6 + (k // cols) * (h + 6)
        sh.paste(im, (x, y)); d.rectangle([x, y, x + 24, y + 14], fill=(126, 0, 41)); d.text((x + 2, y + 1), str(k + 1), fill="white")
    sh.save(out)
    return len(ims)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(2)
    p = sys.argv[1]
    rows = analyse(p)
    prs = Presentation(p)
    _, number_issues, number_off_canvas = _check_slide_numbers(prs)
    warn = 0
    print(f"{'#':>3} {'trans':6} {'entr':>4} {'sal':>3} {'!!':>3} {'comunes':>7} {'formas':>6}")
    for i, tr, ent, ext, nn, shared, nsh, hid in rows:
        flag = ""
        if tr == "Morph" and shared == 0:
            flag = "  ⚠ Morph sin ninguna forma con el mismo nombre que en la anterior"; warn += 1
        elif tr != "Morph" and ent == 0 and i not in (1, len(rows)) and not hid:
            flag = "  ⚠ estática (sin animación ni Morph): revisar"; warn += 1
        print(f"{i:>3} {tr:6} {ent:>4} {ext:>3} {nn:>3} {'' if shared is None else shared:>7} {nsh:>6}{' (oculta)' if hid else ''}{flag}")
    for i, reason in number_issues:
        print(f"  ⚠ diapositiva {i}: paginación {reason}"); warn += 1
    if number_off_canvas:
        print(f"  ℹ paginación fuera del lienzo durante Morph (zoom permitido): {', '.join(map(str, number_off_canvas))}")
    if "--expect" in sys.argv:
        n = int(sys.argv[sys.argv.index("--expect") + 1])
        if n != len(rows):
            print(f"⚠ {len(rows)} diapositivas, se esperaban {n}"); warn += 1
    if "--sheet" in sys.argv:
        out = sys.argv[sys.argv.index("--sheet") + 1]
        print("hoja de miniaturas:", out, sheet(p, out), "páginas")
    print(f"\n{len(rows)} diapositivas · {sum(r[1]=='Morph' for r in rows)} Morph · {sum(r[2] for r in rows)} entradas · {warn} aviso(s)")
