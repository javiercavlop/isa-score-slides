# ISA · SCORE Slides — (c) 2026 Francisco Javier Cavero López y Alejandro García Fernández
# (SCORE Lab, I3US, Universidad de Sevilla; Grupo ISA). Licencia CC BY 4.0 — ver LICENSE.
# Las presentaciones generadas con este kit pueden usarse libremente sin atribución (LICENSE, parte 2).
"""
isa_deck.py — ISA · SCORE Lab presentation kit (python-pptx + raw OOXML).

Builds PowerPoint decks in the ISA / SCORE / Universidad de Sevilla style,
with the transitions and animations the group actually uses:

  * Morph transitions (byObject / byWord / byChar, with Fade fallback) for
    "build" sequences made of duplicated slides.
  * On-slide entrance / exit animations (Wipe, Zoom, Appear, Dissolve,
    Float In, Rise Up, Fly In, Peek In, Blinds, Split, Fade, Checkerboard,
    Circle) triggered on click, with previous, or after previous.

Requirements: python-pptx >= 0.6.21, lxml.  Template: isa-score-template.pptx
(master "isa-v2" + 19 layouts + a hidden seed slide holding the logos).

Quick start
-----------
    from isa_deck import Deck
    d = Deck()                                   # loads the template next to this file
    d.title_slide("My paper title",
                  authors=[("Francisco Javier Cavero", True), "Antonio Ruiz-Cortés"],
                  affiliation="SCORE Lab, I3US, Universidad de Sevilla, Seville, Spain",
                  venue="CAiSE'26", date="10th June 2026")
    s = d.content_slide("Motivation", number=2)
    box = d.card(s, 1.0, 1.6, 4.5, 1.2, "ACCESSIBILITY LAYER", style="blue")
    d.animate(s, box, "wipe", direction="from_left", trigger="click")
    s2 = d.build_step(s)                          # duplicate + morph
    d.spotlight(s2, 0.9, 1.5, 4.7, 1.4)
    d.closing_slide("My paper title", authors=[...], email="me@us.es")
    d.save("out.pptx")

Units: every public geometry argument is in INCHES on a 13.333 x 7.5 in slide
(16:9, 12192000 x 6858000 EMU).
"""
from __future__ import annotations

import copy
import os
import re
from typing import Iterable, Sequence

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "isa-score-template.pptx")


def restore_pptx(b64_path, out_path=None):
    """The design system stores .pptx files as base64 text (`*.pptx.b64.txt`) because its file
    store only serves web media types. Decode one back to a real .pptx."""
    import base64
    out_path = out_path or b64_path[: -len(".b64.txt")]
    with open(b64_path, "rb") as f:
        data = base64.b64decode(b"".join(f.read().split()))
    with open(out_path, "wb") as f:
        f.write(data)
    return out_path


def _resolve_template(path):
    if os.path.exists(path):
        return path
    for cand in (path + ".b64.txt", os.path.join(HERE, os.path.basename(path) + ".b64.txt")):
        if os.path.exists(cand):
            return restore_pptx(cand, path if os.path.isdir(os.path.dirname(path) or ".") else None)
    raise FileNotFoundError(f"{path} (nor {path}.b64.txt) — download kit/isa-score-template.pptx.b64.txt from the design system")

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
P = "{%s}" % NS["p"]

# ----------------------------------------------------------------------------
# Brand tokens (exact values from the source decks)
# ----------------------------------------------------------------------------
COLORS = {
    "crimson": "A8003B",       # theme accent1/accent2/dk2 — titles, primary marks
    "crimson-deep": "7E0029",  # title band, slide numbers, arrows, filled bars
    "crimson-ink": "8B1B3F",   # dark crimson diagram ink / brace
    "signal-red": "C00000",    # agenda circles & frames
    "alert-red": "FF0000",     # spotlight frames (AGF variant), strike/emphasis
    "ink": "141414",           # theme dk1 — body text
    "ink-soft": "222222",
    "grey-text": "7F7F7F",     # body placeholder text (bg1 lumMod 50%)
    "navy": "183548",          # DevOps-loop dark
    "green": "19D65F",         # DevOps-loop light
    "teal-deep": "00504D",     # theme accent3
    "teal": "00A19B",          # theme accent4
    "sky": "4BACC6",           # theme accent5
    "bronze": "967140",        # theme accent6
    "amber": "F9B702",
    "white": "FFFFFF",
}
# draw.io-style pastel cards used for architecture "layers"
CARD_STYLES = {
    "blue":   ("DAE8FC", "A7C6DC", "0070C0"),
    "green":  ("D5E8D4", "97D078", "00B050"),   # source pair; low contrast — keep text large & bold
    "orange": ("FFE6CC", "FFB570", "A8003B"),
    "pink":   ("FDECF4", "F8B2CD", "7E0029"),
    "rose":   ("FCE7EF", "A5003E", "5A0022"),
    "white":  ("FFFFFF", "F8B2CD", "141414"),
    "crimson": ("7E0029", None, "FFFFFF"),
    "outline": (None, "A8003B", "A8003B"),
}

SLIDE_W, SLIDE_H = 12192000, 6858000
NUMBER_BOX = (11725206, 6396335, 466794, 461665)  # logical slide number (EMU)


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def _in(v: float) -> Emu:
    return Emu(int(round(v * 914400)))


# ----------------------------------------------------------------------------
# Animation XML (copied from PowerPoint output found in the source decks)
# ----------------------------------------------------------------------------
def _set_vis(spid, cid, val="visible", delay=0):
    return (f'<p:set><p:cBhvr><p:cTn id="{cid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="{val}"/></p:to></p:set>')


def _anim_effect(spid, cid, flt, dur, transition="in"):
    return (f'<p:animEffect transition="{transition}" filter="{flt}"><p:cBhvr><p:cTn id="{cid}" dur="{dur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')


def _anim_prop(spid, cid, attr, v0, v1, dur, extra="", additive="", fill=' fill="hold"', delay=None, fltfirst=False):
    st = f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>' if delay is not None else ""
    def val(v):
        return f'<p:fltVal val="{v}"/>' if isinstance(v, (int, float)) else f'<p:strVal val="{v}"/>'
    add = f' additive="{additive}"' if additive else ""
    return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr{add}><p:cTn id="{cid}" dur="{dur}"{extra}{fill}>{st}</p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:tavLst><p:tav tm="0"><p:val>{val(v0)}</p:val></p:tav><p:tav tm="100000"><p:val>{val(v1)}</p:val></p:tav></p:tavLst></p:anim>')


_DIR = {"from_left": 8, "from_right": 2, "from_top": 1, "from_bottom": 4}
_WIPE_FILTER = {8: "wipe(left)", 2: "wipe(right)", 1: "wipe(up)", 4: "wipe(down)"}


class _Ids:
    def __init__(self, start):
        self.n = start

    def __call__(self):
        self.n += 1
        return self.n


def _effect_children(effect, cls, sub, spid, ids, dur):
    """Return (presetID, presetSubtype, childTnLst-inner-xml)."""
    e = effect
    if cls == "entr":
        if e == "appear":
            return 1, 0, _set_vis(spid, ids())
        if e == "fade":
            return 10, 0, _set_vis(spid, ids()) + _anim_effect(spid, ids(), "fade", dur)
        if e == "dissolve":
            return 9, 0, _set_vis(spid, ids()) + _anim_effect(spid, ids(), "dissolve", dur)
        if e == "wipe":
            return 22, sub, _set_vis(spid, ids()) + _anim_effect(spid, ids(), _WIPE_FILTER[sub], dur)
        if e == "zoom":  # PowerPoint "Zoom" (subtle): grow from centre + fade
            return 53, 16, (_set_vis(spid, ids()) + _anim_prop(spid, ids(), "ppt_w", 0, "#ppt_w", dur)
                            + _anim_prop(spid, ids(), "ppt_h", 0, "#ppt_h", dur) + _anim_effect(spid, ids(), "fade", dur))
        if e == "float_in":  # Float In (up)
            return 42, 0, (_set_vis(spid, ids()) + _anim_effect(spid, ids(), "fade", dur)
                           + _anim_prop(spid, ids(), "ppt_x", "#ppt_x", "#ppt_x", dur)
                           + _anim_prop(spid, ids(), "ppt_y", "#ppt_y+.1", "#ppt_y", dur))
        if e == "rise_up":  # Rise Up with small overshoot
            d1 = int(dur * 0.9)
            return 37, 0, (_set_vis(spid, ids()) + _anim_effect(spid, ids(), "fade", dur)
                           + _anim_prop(spid, ids(), "ppt_x", "#ppt_x", "#ppt_x", dur)
                           + _anim_prop(spid, ids(), "ppt_y", "#ppt_y+1", "#ppt_y-.03", d1, extra=' decel="100000"')
                           + _anim_prop(spid, ids(), "ppt_y", "#ppt_y-.03", "#ppt_y", dur - d1, extra=' accel="100000"', delay=d1))
        if e == "fly_in":
            x0, y0 = {8: ("0-#ppt_w/2", "#ppt_y"), 2: ("1+#ppt_w/2", "#ppt_y"),
                      1: ("#ppt_x", "0-#ppt_h/2"), 4: ("#ppt_x", "1+#ppt_h/2")}[sub]
            return 2, sub, (_set_vis(spid, ids()) + _anim_prop(spid, ids(), "ppt_x", x0, "#ppt_x", dur, additive="base")
                            + _anim_prop(spid, ids(), "ppt_y", y0, "#ppt_y", dur, additive="base"))
        if e == "peek_in":
            if sub in (8, 2):
                attr, v0 = "ppt_x", ("#ppt_x-#ppt_w*1.125000" if sub == 8 else "#ppt_x+#ppt_w*1.125000")
                flt = "wipe(right)" if sub == 8 else "wipe(left)"
            else:
                attr, v0 = "ppt_y", ("#ppt_y+#ppt_h*1.125000" if sub == 4 else "#ppt_y-#ppt_h*1.125000")
                flt = "wipe(up)" if sub == 4 else "wipe(down)"
            return 12, sub, (_set_vis(spid, ids()) + _anim_prop(spid, ids(), attr, v0, "#" + attr, dur, additive="base", fill="")
                             + _anim_effect(spid, ids(), flt, dur))
        if e == "blinds":
            return 3, 10, _set_vis(spid, ids()) + _anim_effect(spid, ids(), "blinds(horizontal)", dur)
        if e == "split":
            return 16, 21, _set_vis(spid, ids()) + _anim_effect(spid, ids(), "barn(inVertical)", dur)
        if e == "checkerboard":
            return 5, 10, _set_vis(spid, ids()) + _anim_effect(spid, ids(), "checkerboard(across)", dur)
        if e == "circle":
            return 6, 32, _set_vis(spid, ids()) + _anim_effect(spid, ids(), "circle(out)", dur)
    else:  # exit
        last = max(dur - 1, 0)
        if e in ("disappear", "appear"):
            return 1, 0, _set_vis(spid, ids(), "hidden")
        if e == "dissolve":
            return 9, 0, _anim_effect(spid, ids(), "dissolve", dur, "out") + _set_vis(spid, ids(), "hidden", last)
        if e == "fade":
            return 10, 0, _anim_effect(spid, ids(), "fade", dur, "out") + _set_vis(spid, ids(), "hidden", last)
        if e == "wipe":
            return 22, sub, _anim_effect(spid, ids(), _WIPE_FILTER[sub], dur, "out") + _set_vis(spid, ids(), "hidden", last)
        if e == "zoom":
            return 53, 32, (_anim_prop(spid, ids(), "ppt_w", "ppt_w", 0, dur, fill="")
                            + _anim_prop(spid, ids(), "ppt_h", "ppt_h", 0, dur, fill="")
                            + _anim_effect(spid, ids(), "fade", dur, "out") + _set_vis(spid, ids(), "hidden", last))
    raise ValueError(f"unsupported effect {effect!r} for class {cls!r}")


ENTRANCE_EFFECTS = ["appear", "fade", "dissolve", "wipe", "zoom", "float_in", "rise_up", "fly_in",
                    "peek_in", "blinds", "split", "checkerboard", "circle"]
EXIT_EFFECTS = ["disappear", "dissolve", "fade", "wipe", "zoom"]
DEFAULT_DUR = {"rise_up": 1000, "float_in": 1000, "circle": 2000}


# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def _remove_children(el, tags):
    for t in tags:
        for c in el.findall(t):
            el.remove(c)


def set_transition(slide, kind="morph", option="byObject", speed="slow", dur_ms=None, advance_after_ms=None):
    """Set the slide's transition.

    kind: "morph" | "fade" | "none" | "push" | "wipe" | "split"
    option (morph): "byObject" (default, used in ~98% of source slides) | "byWord" | "byChar"
    dur_ms: explicit duration (p14:dur). Source decks use the 'slow' default, 2000 ms for
            dramatic morphs, and 50 ms for an instant 'cut' morph that keeps objects in place.
    advance_after_ms: auto-advance (advTm), e.g. 1500 for timed sequences.
    """
    sld = slide._element
    for c in list(sld):
        if c.tag in (P + "transition",) or c.tag.endswith("}AlternateContent"):
            if c.tag.endswith("}AlternateContent") and b"transition" not in etree.tostring(c):
                continue
            sld.remove(c)
    if kind == "none":
        return
    base = []
    if speed:
        base.append(f'spd="{speed}"')
    if advance_after_ms is not None:
        base.append(f'advTm="{advance_after_ms}"')
    p14 = 'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"'
    a_plain = " ".join(base)
    a = " ".join(base + ([f'{p14} p14:dur="{dur_ms}"'] if dur_ms is not None else []))
    if kind == "morph":
        xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
               f'xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" xmlns:p="{NS["p"]}">'
               f'<mc:Choice Requires="p159">'
               f'<p:transition {a}><p159:morph option="{option}"/></p:transition></mc:Choice>'
               f'<mc:Fallback><p:transition {a_plain}><p:fade/></p:transition></mc:Fallback>'
               f'</mc:AlternateContent>')
    else:
        inner = {"fade": "<p:fade/>", "push": '<p:push dir="u"/>', "wipe": "<p:wipe/>", "split": '<p:split orient="vert"/>'}[kind]
        xml = f'<p:transition xmlns:p="{NS["p"]}" {a}>{inner}</p:transition>'
    el = etree.fromstring(xml)
    # insert after clrMapOvr (before timing / extLst)
    anchor = sld.find(P + "clrMapOvr")
    if anchor is None:
        anchor = sld.find(P + "cSld")
    anchor.addnext(el)


def _shape_id(shape):
    return shape.shape_id if hasattr(shape, "shape_id") else int(shape)


def apply_animations(slide, steps: Sequence[dict]):
    """Write the slide's animation timeline.

    steps: ordered list of dicts:
        shape      python-pptx shape (or shape id)
        effect     one of ENTRANCE_EFFECTS / EXIT_EFFECTS
        cls        "entr" (default) | "exit"
        direction  "from_left" | "from_right" | "from_top" | "from_bottom" (wipe, fly_in, peek_in)
        trigger    "click" | "with" | "after"   (default "click")
        delay      ms (default 0)
        dur        ms (default 500; 1000 for rise_up/float_in, 2000 for circle)

    Replaces any existing timeline. Same semantics as PowerPoint's Animation Pane:
    a 'click' starts a new click group; 'with' runs alongside the previous effect;
    'after' starts when the previous effect ends.
    """
    sld = slide._element
    for c in sld.findall(P + "timing"):
        sld.remove(c)
    if not steps:
        return
    ids = _Ids(2)
    groups = []  # list of (autostart, [subpars]) ; subpar = (delay, [effect xml], length)
    for i, st in enumerate(steps):
        trig = st.get("trigger", "click")
        if trig == "click" or not groups:
            groups.append({"auto": trig != "click", "subs": []})
            g = groups[-1]
            g["subs"].append({"delay": 0, "fx": [], "len": 0})
        else:
            g = groups[-1]
            if trig == "after":
                prev = g["subs"][-1]
                g["subs"].append({"delay": prev["delay"] + prev["len"], "fx": [], "len": 0})
        g = groups[-1]
        sub = g["subs"][-1]
        st["_node"] = {"click": "clickEffect", "with": "withEffect", "after": "afterEffect"}[trig]
        sub["fx"].append(st)
        eff = st["effect"]
        dur = st.get("dur", DEFAULT_DUR.get(eff, 500))
        st["_dur"] = dur
        sub["len"] = max(sub["len"], st.get("delay", 0) + dur)

    grp_ids = {}
    bld = []
    out = []
    for g in groups:
        outer_id = ids()
        conds = '<p:cond delay="indefinite"/>' + ('<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>' if g["auto"] else "")
        subxml = []
        for sub in g["subs"]:
            sid = ids()
            fxx = []
            for st in sub["fx"]:
                cls = st.get("cls", "entr")
                eff = st["effect"]
                d = _DIR.get(st.get("direction", "from_left"), 8)
                spid = _shape_id(st["shape"])
                cid = ids()
                pid, psub, inner = _effect_children(eff, cls, d, spid, ids, st["_dur"])
                key = (spid, cls)
                if key not in grp_ids:
                    grp_ids[key] = len([k for k in grp_ids if k[0] == spid])
                    shp = st["shape"]
                    is_sp = hasattr(shp, "_element") and shp._element.tag == P + "sp"
                    if is_sp:
                        bld.append(f'<p:bldP spid="{spid}" grpId="{grp_ids[key]}" animBg="1"/>')
                grp = grp_ids[key]
                fxx.append(f'<p:par><p:cTn id="{cid}" presetID="{pid}" presetClass="{cls}" presetSubtype="{psub}" fill="hold" '
                           f'grpId="{grp}" nodeType="{st["_node"]}"><p:stCondLst><p:cond delay="{st.get("delay", 0)}"/></p:stCondLst>'
                           f'<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>')
            subxml.append(f'<p:par><p:cTn id="{sid}" fill="hold"><p:stCondLst><p:cond delay="{sub["delay"]}"/></p:stCondLst>'
                          f'<p:childTnLst>{"".join(fxx)}</p:childTnLst></p:cTn></p:par>')
        out.append(f'<p:par><p:cTn id="{outer_id}" fill="hold"><p:stCondLst>{conds}</p:stCondLst>'
                   f'<p:childTnLst>{"".join(subxml)}</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{NS["p"]}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(out)}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst>'
           + (f'<p:bldLst>{"".join(bld)}</p:bldLst>' if bld else "") + '</p:timing>')
    el = etree.fromstring(xml)
    # timing goes after transition/AlternateContent, before extLst
    ext = sld.find(P + "extLst")
    if ext is not None:
        ext.addprevious(el)
    else:
        sld.append(el)


def duplicate_slide(prs, src, position=None):
    """Clone `src` (shapes, pictures, background) right after it (or at `position`).

    Transition and animations are NOT copied — set them on the clone explicitly.
    Shape names and ids are preserved, so Morph matches every object automatically.
    """
    new = prs.slides.add_slide(prs.slide_layouts[0])
    _retarget_layout(new, _layout_part_of(src))
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag not in (P + "nvGrpSpPr", P + "grpSpPr"):
            tree.remove(el)
    rid_map = {}
    _keep_rids(new.part, src.part, lambda rel: None if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE) else
               ((rel.target_ref, True) if rel.is_external else (rel.target_part, False)))
    for el in src.shapes._spTree:
        if el.tag in (P + "nvGrpSpPr", P + "grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    for el in tree.iter():
        for attr in ("{%s}embed" % NS["r"], "{%s}link" % NS["r"], "{%s}id" % NS["r"]):
            v = el.get(attr)
            if v in rid_map:
                el.set(attr, rid_map[v])
    # background + master-shape visibility
    sbg = src._element.find(P + "cSld/" + P + "bg")
    if sbg is not None:
        new._element.find(P + "cSld").insert(0, copy.deepcopy(sbg))
    if src._element.get("showMasterSp") is not None:
        new._element.set("showMasterSp", src._element.get("showMasterSp"))
    # reorder
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new_id = ids[-1]
    lst.remove(new_id)
    if position is None:
        position = _index_of(prs, src) + 1
    lst.insert(min(position, len(lst)), new_id)
    return new



# ----------------------------------------------------------------------------
# Cross-presentation slide import (keeps media, Morph transitions and timings)
# ----------------------------------------------------------------------------
_SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE,
              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml",
              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags",
              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"}
_R_ATTRS = None


def _compress_image(blob, content_type, max_bytes=500_000, max_w=1440):
    if len(blob) <= max_bytes or content_type not in ("image/png", "image/jpeg"):
        return blob
    try:
        from io import BytesIO
        from PIL import Image
        im = Image.open(BytesIO(blob))
        if im.width > max_w:
            im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
        out = BytesIO()
        if content_type == "image/png":
            im.save(out, "PNG", optimize=True)
            if out.tell() > max_bytes:  # flat artwork (logos, icons): palette PNG keeps alpha
                out = BytesIO()
                im.convert("RGBA").quantize(256, method=Image.FASTOCTREE).save(out, "PNG", optimize=True)
        else:
            im.convert("RGB").save(out, "JPEG", quality=88, optimize=True)
        return out.getvalue() if out.tell() < len(blob) else blob
    except Exception:
        return blob


def _copy_part(src_part, dest_package, memo, compress=True):
    """Copy a part (and, recursively, the parts it relates to) into another package."""
    from pptx.opc.package import Part, _Relationship
    from pptx.opc.packuri import PackURI
    key = id(src_part)
    if key in memo:
        return memo[key]
    pn = str(src_part.partname)
    m = re.match(r"^(.*?)(\d+)(\.\w+)$", pn)
    tmpl = (m.group(1) + "%d" + m.group(3)) if m else pn.rsplit(".", 1)[0] + "%d." + pn.rsplit(".", 1)[1]
    # unique partname: parts copied in this recursion are not reachable yet, so reserve names explicitly
    reserved = dest_package.__dict__.setdefault("_isa_reserved", set())
    used = {str(p.partname) for p in dest_package.iter_parts()} | reserved
    n = 1
    while tmpl % n in used:
        n += 1
    new_pn = tmpl % n
    reserved.add(new_pn)
    blob = src_part.blob
    if compress:
        blob = _compress_image(blob, src_part.content_type)
    new = Part(PackURI(new_pn), src_part.content_type, dest_package, blob)
    memo[key] = new
    for rId, rel in src_part.rels.items():
        if rel.reltype in _SKIP_RELS and rel.reltype != RT.SLIDE_LAYOUT:  # masters keep their layouts
            continue
        if rel.is_external:
            tgt, mode = rel.target_ref, "External"
        else:
            tgt, mode = _copy_part(rel.target_part, dest_package, memo, compress), "Internal"
        new.rels._rels[rId] = _Relationship(new.partname.baseURI, rId, rel.reltype, mode, tgt)
    return new


def _xml_of(part):
    el = getattr(part, "_element", None)
    return el if el is not None else etree.fromstring(part.blob)


def _master_key(master_part):
    """Content key of a slide master + its theme + its layouts (ids/rIds ignored)."""
    import hashlib
    h = hashlib.md5()
    root = copy.deepcopy(_xml_of(master_part))
    lst = root.find(P + "sldLayoutIdLst")
    if lst is not None:
        root.remove(lst)
    h.update(etree.tostring(root, method="c14n"))
    for rel in sorted(master_part.rels.values(), key=lambda r: r.reltype):
        if rel.is_external:
            continue
        if rel.reltype == RT.THEME:
            h.update(etree.tostring(_xml_of(rel.target_part), method="c14n"))
    lays = []
    for r in master_part.rels.values():
        if not r.is_external and r.reltype == RT.SLIDE_LAYOUT:
            lays.append(etree.tostring(_xml_of(r.target_part), method="c14n"))
    for b in sorted(lays):  # layout contents too: same names are not enough
        h.update(b)
    return h.hexdigest()


def _layout_name(layout_part):
    c = _xml_of(layout_part).find(P + "cSld")
    return c.get("name", "") if c is not None else ""


def _layout_part_of(slide):
    return slide.part.part_related_by(RT.SLIDE_LAYOUT)


def _master_part_of(layout_part):
    return layout_part.part_related_by(RT.SLIDE_MASTER)


def _dest_masters(dest_prs):
    cache = getattr(dest_prs.part, "_isa_master_cache", None)
    if cache is None:
        cache = {}
        for rel in dest_prs.part.rels.values():
            if rel.reltype == RT.SLIDE_MASTER:
                cache[_master_key(rel.target_part)] = rel.target_part
        dest_prs.part._isa_master_cache = cache
    return cache


def _all_master_layout_ids(dest_prs):
    ids = [int(x.get("id")) for x in dest_prs.part._element.iter(P + "sldMasterId")]
    for rel in dest_prs.part.rels.values():
        if rel.reltype == RT.SLIDE_MASTER:
            ids += [int(x.get("id")) for x in _xml_of(rel.target_part).iter(P + "sldLayoutId")]
    return ids


def _resolve_layout(dest_prs, src_layout_part, memo, compress):
    """Destination layout part equivalent to `src_layout_part`: the same-named layout of an
    identical master if the destination already has one (e.g. the isa-v2 master of the template),
    otherwise the source master is imported (theme, all its layouts, its images) and registered."""
    src_master = _master_part_of(src_layout_part)
    key = _master_key(src_master)
    masters = _dest_masters(dest_prs)
    name = _layout_name(src_layout_part)
    if key in masters:
        dm = masters[key]
        for rel in dm.rels.values():
            if rel.reltype == RT.SLIDE_LAYOUT and _layout_name(rel.target_part) == name:
                return rel.target_part
    # import the master (recursively copies theme, layouts and their media)
    new_master = _copy_part(src_master, dest_prs.part.package, memo, compress)
    root = etree.fromstring(new_master.blob)
    nxt = max(_all_master_layout_ids(dest_prs) + [2147483648]) + 1
    for lid in root.iter(P + "sldLayoutId"):
        lid.set("id", str(nxt)); nxt += 1
    new_master._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    rId = dest_prs.part.relate_to(new_master, RT.SLIDE_MASTER)
    lst = dest_prs.part._element.find(P + "sldMasterIdLst")
    el = etree.SubElement(lst, P + "sldMasterId")
    el.set("id", str(nxt)); el.set("{%s}id" % NS["r"], rId)
    masters[key] = new_master
    if id(src_layout_part) not in memo:  # orphan layout (used by a slide, not listed by its master)
        lay = _copy_part(src_layout_part, dest_prs.part.package, memo, compress)
        lrid = new_master.relate_to(lay, RT.SLIDE_LAYOUT)
        root = etree.fromstring(new_master.blob)
        lst2 = root.find(P + "sldLayoutIdLst")
        if lst2 is None:
            lst2 = etree.SubElement(root, P + "sldLayoutIdLst")
        e2 = etree.SubElement(lst2, P + "sldLayoutId")
        e2.set("id", str(max(_all_master_layout_ids(dest_prs) + [nxt]) + 1)); e2.set("{%s}id" % NS["r"], lrid)
        new_master._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return memo[id(src_layout_part)]


def _retarget_layout(slide, layout_part):
    for rel in slide.part.rels.values():
        if rel.reltype == RT.SLIDE_LAYOUT:
            rel._target = layout_part
            return


def _keep_rids(new_part, src_part, target_for):
    """Give `new_part` the SAME relationship ids as `src_part` (SmartArt data parts and other
    parts refer to the slide's rIds internally, e.g. dsp:dataModelExt/@relId). `target_for(rel)`
    returns (target, is_external) or None to skip. The pre-existing layout rel is renamed if needed."""
    from pptx.opc.package import _Relationship
    wanted = {rId for rId, rel in src_part.rels.items() if target_for(rel) is not None}
    rels = new_part.rels._rels
    for rId in list(rels):
        if rId in wanted:
            rel = rels.pop(rId)
            n = 1
            while f"rId{n}" in wanted or f"rId{n}" in rels:
                n += 1
            rels[f"rId{n}"] = _Relationship(new_part.partname.baseURI, f"rId{n}", rel.reltype,
                                            "External" if rel.is_external else "Internal",
                                            rel.target_ref if rel.is_external else rel.target_part)
    for rId, rel in src_part.rels.items():
        t = target_for(rel)
        if t is None:
            continue
        tgt, ext = t
        rels[rId] = _Relationship(new_part.partname.baseURI, rId, rel.reltype, "External" if ext else "Internal", tgt)


def import_slide(dest_prs, src_slide, memo=None, compress=False, x_offset=0, layout_fallback="En blanco", position=None):
    """Copy `src_slide` from another presentation into `dest_prs` (appended, or at `position`).

    Everything travels: shapes, groups, pictures/SVG, SmartArt, background, hidden flag, the Morph /
    other transition (incl. its mc:AlternateContent) and the full animation timeline. The layout is
    kept faithful: if the destination has an identical master (same theme, styles and layouts) the
    same-named layout is used; otherwise the source master, theme and layouts are imported too, so
    theme colours, fonts and placeholder styles render exactly as in the original deck.
    x_offset (EMU) shifts every top-level shape — use 1524000 to centre a 4:3 slide on 16:9."""
    memo = {} if memo is None else memo
    target_layout = _resolve_layout(dest_prs, _layout_part_of(src_slide), memo, compress)
    if x_offset:  # 4:3 → 16:9: give inherited placeholders explicit positions so they move too
        for ph in src_slide.placeholders:
            try:
                if ph._element.find(".//" + "{%s}xfrm" % NS["a"]) is None:
                    l, t, w, h = ph.left, ph.top, ph.width, ph.height
                    if None not in (l, t, w, h):
                        ph.left, ph.top, ph.width, ph.height = l, t, w, h
            except Exception:
                pass
    new = dest_prs.slides.add_slide(dest_prs.slide_layouts[0])
    _retarget_layout(new, target_layout)
    sld_new, sld_src = new._element, src_slide._element
    def _tgt(rel):
        if rel.reltype in _SKIP_RELS:
            return None
        if rel.is_external:
            return rel.target_ref, True
        return _copy_part(rel.target_part, dest_prs.part.package, memo, compress), False
    _keep_rids(new.part, src_slide.part, _tgt)
    rid_map = {}
    # copy cSld contents IN PLACE (python-pptx caches the cSld/spTree objects), then everything after it
    csld_new, csld_src = sld_new.find(P + "cSld"), sld_src.find(P + "cSld")
    tree_new, tree_src = csld_new.find(P + "spTree"), csld_src.find(P + "spTree")
    for c in list(tree_new):
        tree_new.remove(c)
    for c in tree_src:
        tree_new.append(copy.deepcopy(c))
    for c in list(csld_new):
        if c is not tree_new:
            csld_new.remove(c)
    for c in csld_src:
        if c.tag == P + "bg":
            tree_new.addprevious(copy.deepcopy(c))
        elif c.tag != P + "spTree":
            csld_new.append(copy.deepcopy(c))
    for k, v in csld_src.attrib.items():
        csld_new.set(k, v)
    for c in list(sld_new):
        if c is not csld_new:
            sld_new.remove(c)
    for c in sld_src:
        if c.tag != P + "cSld":
            sld_new.append(copy.deepcopy(c))
    for k, v in sld_src.attrib.items():
        sld_new.set(k, v)
    # drop customer data / tags (not copied)
    for cd in sld_new.findall(".//" + P + "custDataLst"):
        cd.getparent().remove(cd)
    rns = "{%s}" % NS["r"]
    for el in sld_new.iter():
        for attr, v in list(el.attrib.items()):
            if attr.startswith(rns) and v in rid_map:
                el.set(attr, rid_map[v])
    if x_offset:
        tree = sld_new.find(P + "cSld/" + P + "spTree")
        for el in tree:
            off = el.find("./*/" + "{%s}xfrm/{%s}off" % (NS["a"], NS["a"]))
            if off is None:
                off = el.find("./" + P + "xfrm/{%s}off" % NS["a"])
            if off is not None:
                off.set("x", str(int(off.get("x")) + x_offset))
    if position is not None:
        lst = dest_prs.slides._sldIdLst
        sid = list(lst)[-1]
        lst.remove(sid)
        lst.insert(position, sid)
    return new


def import_slides(dest_prs, src, slide_numbers=None, include_hidden=False, **kw):
    """Import several slides (1-based numbers; default all) from a path or Presentation."""
    sp = Presentation(src) if isinstance(src, str) else src
    memo = {}
    out = []
    for i, s in enumerate(sp.slides, 1):
        if slide_numbers is not None and i not in slide_numbers:
            continue
        if not include_hidden and s._element.get("show") == "0":
            continue
        out.append(import_slide(dest_prs, s, memo=memo, **kw))
    return out


def replace_text(slides, mapping: dict):
    """Replace text in every run of `slides` (str → str). A key that spans several runs of one
    paragraph is also matched: the paragraph keeps the first run's formatting."""
    if not mapping:
        return
    for s in slides:
        for p in s._element.iter("{%s}p" % NS["a"]):
            runs = p.findall("{%s}r" % NS["a"])
            for r in runs:
                t = r.find("{%s}t" % NS["a"])
                if t is not None and t.text:
                    for k, v in mapping.items():
                        if k in t.text:
                            t.text = t.text.replace(k, v)
            full = "".join((r.find("{%s}t" % NS["a"]).text or "") for r in runs if r.find("{%s}t" % NS["a"]) is not None)
            for k, v in mapping.items():
                if k in full and len(runs) > 1 and not any(k in (r.find("{%s}t" % NS["a"]).text or "") for r in runs):
                    ts = [r.find("{%s}t" % NS["a"]) for r in runs]
                    ts[0].text = full.replace(k, v)
                    for t in ts[1:]:
                        t.text = ""


def set_logical_number(slides, number):
    """Set the bottom-right logical slide number text box (digits or 'n/N') on each slide."""
    for s in slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.left is not None and sh.left > 10_800_000 and sh.top > 6_000_000:
                txt = sh.text_frame.text.strip()
                if re.fullmatch(r"\d+(/\d+)?", txt):
                    runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
                    if runs:
                        runs[0].text = str(number)
                        for r in runs[1:]:
                            r.text = ""


def _index_of(prs, slide):
    for i, s in enumerate(prs.slides):
        if s.slide_id == slide.slide_id:
            return i
    raise ValueError("slide not in presentation")


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
        if sh.shape_type == 6:  # group
            for c in sh.shapes:
                if c.name == name:
                    return c
    raise KeyError(name)


# ----------------------------------------------------------------------------
# Deck
# ----------------------------------------------------------------------------
LAYOUT = {
    "title": "Diapositiva de título",
    "content": "Título y objetos",
    "title_only": "Sólo el título",
    "blank": "En blanco",
    "two": "Dos objetos",
    "compare": "Comparación",
    "section": "Encabezado de sección",
    "central": "Central con titulo",
    "big_image": "Titulo y imagen grande",
    "full_image": "Solo imagen",
}


class Deck:
    def __init__(self, template: str = DEFAULT_TEMPLATE, title_font: str | None = None, body_font: str = "Roboto"):
        self.prs = Presentation(_resolve_template(template))
        # generated decks carry no attribution or metadata from the design system (LICENSE, part 2)
        cp = self.prs.core_properties
        for k in ("title", "author", "subject", "keywords", "comments", "last_modified_by", "category"):
            try:
                setattr(cp, k, "")
            except Exception:
                pass
        self.body_font = body_font
        self.title_font = title_font  # None → theme heading font (Arial Narrow, as in the master)
        self._assets = {}
        self._seed = None
        self._seeds = {}
        for s in self.prs.slides:
            nm = s._element.find(P + "cSld").get("name", "")
            if nm.startswith("ISA seed: "):
                self._seeds[nm[10:]] = s
            names = [sh.name for sh in s.shapes]
            if any(n.startswith("asset:") for n in names):
                self._seed = s
                for sh in s.shapes:
                    if sh.name.startswith("asset:"):
                        self._assets[sh.name[6:]] = sh
        self._name_counter = 0

    # ---------- layouts -------------------------------------------------------
    def layout(self, key):
        name = LAYOUT.get(key, key)
        for l in self.prs.slide_layouts:
            if l.name == name:
                return l
        raise KeyError(name)

    def new_slide(self, layout="title_only"):
        s = self.prs.slides.add_slide(self.layout(layout))
        return s

    # ---------- assets --------------------------------------------------------
    def logo(self, slide, name, x, y, w=None, h=None, shape_name=None):
        """Place a brand asset (isa-logo, score-logo, us-logo, us-seal, i3us-logo,
        junta-andalucia, funding-ministerio-nextgen, score-us-lockup) at inches."""
        src = self._assets[name]
        rid = None
        blip = src._element.find(".//a:blip", NS)
        old = blip.get("{%s}embed" % NS["r"])
        part = self._seed.part.related_part(old)
        rid = slide.part.relate_to(part, RT.IMAGE)
        el = copy.deepcopy(src._element)
        el.find(".//a:blip", NS).set("{%s}embed" % NS["r"], rid)
        ratio = src.height / src.width
        if w is None and h is None:
            w = src.width / 914400
        if w is None:
            w = h / ratio
        if h is None:
            h = w * ratio
        off = el.find(".//a:xfrm/a:off", NS)
        ext = el.find(".//a:xfrm/a:ext", NS)
        off.set("x", str(int(_in(x)))); off.set("y", str(int(_in(y))))
        ext.set("cx", str(int(_in(w)))); ext.set("cy", str(int(_in(h))))
        nv = el.find(".//p:cNvPr", NS)
        nv.set("id", str(self._next_id(slide)))
        nv.set("name", shape_name or f"Logo {name}")
        slide.shapes._spTree.append(el)
        return slide.shapes[-1]

    def _next_id(self, slide):
        ids = [int(e.get("id")) for e in slide._element.iter() if e.tag.endswith("}cNvPr") and e.get("id", "").isdigit()]
        return max(ids + [1]) + 1

    # ---------- text helpers --------------------------------------------------
    def _style_run(self, run, size=None, bold=None, color=None, font=None, italic=None, underline=None):
        f = run.font
        if size: f.size = Pt(size)
        if bold is not None: f.bold = bold
        if italic is not None: f.italic = italic
        if underline is not None: f.underline = underline
        if color: f.color.rgb = _rgb(COLORS.get(color, color))
        if font: f.name = font

    def text(self, slide, x, y, w, h, text, size=20, color="ink", bold=False, align="left", font=None,
             anchor="top", name=None, italic=False):
        """Text box. `text` may be a str, or a list of paragraphs; each paragraph a str or a list
        of (str, {opts}) runs. Opts: size, bold, color, italic, underline, font."""
        tb = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
        if name: tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        paras = text if isinstance(text, list) else [text]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[align]
            runs = para if isinstance(para, list) else [(para, {})]
            for rt, opts in runs:
                r = p.add_run()
                r.text = rt
                o = dict(size=size, bold=bold, color=color, font=font or self.body_font, italic=italic)
                o.update(opts)
                self._style_run(r, **o)
        return tb

    # ---------- slide frames --------------------------------------------------
    # ---------- title & closing: the group's REAL slides with {{placeholders}} ------------
    @staticmethod
    def _split_date(date):
        m = re.match(r"^\s*(\d{1,2})(st|nd|rd|th)?\s+(.*)$", date or "")
        if m:
            return m.group(1), m.group(2) or "", m.group(3)
        return date or "", "", ""

    @staticmethod
    def _split_authors(authors, presenter=None, coauthors=None):
        if presenter is not None:
            return presenter, (coauthors if isinstance(coauthors, str) else ", ".join(coauthors or []))
        al = list(authors or [])
        names = [a if isinstance(a, str) else a[0] for a in al]
        flagged = [i for i, a in enumerate(al) if not isinstance(a, str) and a[1]]
        k = flagged[0] if flagged else 0
        pres = names[k]
        rest = [n for i, n in enumerate(names) if i != k]
        if k != 0:  # keep the original order: presenter underlined where it was
            rest = names[:k] + names[k + 1:]
        return pres, ", ".join(rest)

    def _fill_seed(self, kind, values):
        seed = self._seeds[kind]
        s = duplicate_slide(self.prs, seed, position=len(self.prs.slides._sldIdLst))
        s._element.find(P + "cSld").attrib.pop("name", None)
        replace_text([s], {"{{%s}}" % k: v for k, v in values.items()})
        # ', {{COAUTHORS}}' with no co-authors → drop the comma
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text == ", ":
                            r.text = ""
        return s

    def title_slide(self, title, authors=None, affiliation="SCORE Lab, I3US, Universidad de Sevilla, Seville, Spain",
                    venue="", date="", presenter=None, coauthors=None, **_ignored):
        """Opening slide = the group's real title slide (CAiSE'26) with its placeholders filled:
        {{TITLE}} · {{PRESENTER}} (underlined) · {{COAUTHORS}} · {{AFFILIATION}} · {{VENUE}} ·
        {{DAY}}{{DAY_SUFFIX}} {{MONTH_YEAR}} (date like '10th June 2026'; the suffix is superscript).
        authors: list of str or (name, is_presenter); or pass presenter= and coauthors=."""
        pres, co = self._split_authors(authors, presenter, coauthors)
        day, suf, rest = self._split_date(date)
        return self._fill_seed("title", {"TITLE": title, "PRESENTER": pres, "COAUTHORS": co, "AFFILIATION": affiliation,
                                         "VENUE": venue, "DAY": day, "DAY_SUFFIX": suf, "MONTH_YEAR": rest})

    def closing_slide(self, title, authors=None, email="", affiliation="SCORE Lab, I3US, Universidad de Sevilla, Seville, Spain",
                      venue="", date="", qr_image=None, funding=True, presenter=None, coauthors=None):
        """Closing slide = the group's real closing slide (CAiSE'26): same logos, band and title,
        authors + ({{EMAIL}}), QR at the left and the funding logos bottom-right.
        qr_image: path to the QR PNG; without it the slide keeps a 'QR' placeholder picture you can
        swap in PowerPoint (right click → Change Picture)."""
        pres, co = self._split_authors(authors, presenter, coauthors)
        day, suf, rest = self._split_date(date)
        s = self._fill_seed("closing", {"TITLE": title, "PRESENTER": pres, "COAUTHORS": co, "EMAIL": email,
                                        "AFFILIATION": affiliation, "VENUE": venue, "DAY": day, "DAY_SUFFIX": suf,
                                        "MONTH_YEAR": rest})
        for sh in list(s.shapes):
            if sh.name == "QR placeholder" and qr_image:
                img, _ = s.part.get_or_add_image_part(open(qr_image, "rb") if isinstance(qr_image, str) else qr_image)
                rid = s.part.relate_to(img, RT.IMAGE)
                sh._element.find(".//a:blip", NS).set("{%s}embed" % NS["r"], rid)
                sh.name = "QR"
            if sh.name == "Funding logos" and not funding:
                sh._element.getparent().remove(sh._element)
        return s

    def content_slide(self, title, number=None, layout="title_only", title_size=None, title_color=None):
        """Standard content frame: centred bold title (theme heading font, 28 pt, crimson #A8003B)
        and a logical slide number bottom-right (24 pt, #7E0029). Build sequences repeat the SAME number."""
        s = self.new_slide(layout)
        t = s.shapes.title
        t.text_frame.text = title
        t.name = "Title"
        r = t.text_frame.paragraphs[0].runs[0]
        if self.title_font: r.font.name = self.title_font
        if title_size: r.font.size = Pt(title_size)
        if title_color: r.font.color.rgb = _rgb(COLORS.get(title_color, title_color))
        # drop empty body placeholders
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        if number is not None:
            self.slide_number(s, number)
        return s

    def slide_number(self, slide, number, size=24):
        """Manual logical number (the group numbers *ideas*, not physical slides: every slide of a
        morph build carries the same number). Use 'n/N' strings for PhD-forum style."""
        x, y, w, h = NUMBER_BOX
        tb = slide.shapes.add_textbox(Emu(x - 300000), Emu(y), Emu(w + 300000), Emu(h))
        tb.name = "Slide number"
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = str(number)
        self._style_run(r, size=size, color="crimson-deep", font=self.body_font)
        return tb

    def citation(self, slide, text, size=9):
        """Reference footer bottom-left (APA, venue in italics allowed via list of runs)."""
        return self.text(slide, 0.08, 6.82, 11.0, 0.5, text, size=size, color="ink-soft", name="Citation", anchor="bottom")

    def section_divider(self, text, number=None, accent_words=None):
        """Big centred statement on white (e.g. 'Pricing-driven Development and Operation'),
        crimson 32 pt; optional words highlighted in deep crimson."""
        s = self.new_slide("blank")
        runs = []
        for w in text.split(" "):
            col = "crimson-deep" if accent_words and w.strip(",.") in accent_words else "crimson"
            runs.append((w + " ", {"color": col}))
        self.text(s, 0.8, 2.9, 11.7, 1.4, [runs], size=36, align="center", anchor="middle", name="Statement")
        if number is not None:
            self.slide_number(s, number)
        return s

    # ---------- components ----------------------------------------------------
    def card(self, slide, x, y, w, h, text="", style="blue", size=18, bold=False, radius=0.0585, name=None,
             align="center", font=None, line_w=1.5):
        """Rounded 'layer' card in the draw.io pastel palette (blue / green / orange / pink / rose /
        white / crimson / outline)."""
        fill, line, ink = CARD_STYLES[style]
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _in(x), _in(y), _in(w), _in(h))
        sh.adjustments[0] = radius
        sh.shadow.inherit = False
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(fill)
        else: sh.fill.background()
        if line: sh.line.color.rgb = _rgb(line); sh.line.width = Pt(line_w)
        else: sh.line.fill.background()
        if name: sh.name = name
        tf = sh.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if text:
            paras = text if isinstance(text, list) else [text]
            for i, ptxt in enumerate(paras):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
                r = p.add_run(); r.text = ptxt
                self._style_run(r, size=size, bold=bold if i == 0 else False, color=ink, font=font or self.body_font)
        return sh

    def spotlight(self, slide, x, y, w, h, color="crimson-deep", width_pt=2.25, name="!!Spotlight", shape="rect"):
        """Hollow frame used to point at part of a diagram. Keep the SAME name on every slide of a
        build so Morph glides it from target to target."""
        geom = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL}[shape]
        sh = slide.shapes.add_shape(geom, _in(x), _in(y), _in(w), _in(h))
        sh.fill.background()
        sh.line.color.rgb = _rgb(COLORS.get(color, color)); sh.line.width = Pt(width_pt)
        sh.shadow.inherit = False
        sh.name = name
        return sh

    def focus_overlay(self, slide, transparency=0.45, name="!!Overlay"):
        """Full-slide grey veil (#7F7F7F) to push the diagram back while a card is shown on top."""
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(SLIDE_W), Emu(SLIDE_H))
        sh.fill.solid(); sh.fill.fore_color.rgb = _rgb("7F7F7F")
        sh.line.fill.background(); sh.shadow.inherit = False
        sh.name = name
        # alpha
        srgb = sh.fill._xPr.find(".//a:srgbClr", NS)
        a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int((1 - transparency) * 100000)))
        return sh

    def quote_card(self, slide, x, y, w, h, quote_runs, name="Quote card"):
        """White card with crimson frame holding an italic definition / quote."""
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _in(x), _in(y), _in(w), _in(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = _rgb("FFFFFF")
        sh.line.color.rgb = _rgb("C00000"); sh.line.width = Pt(2.25)
        sh.shadow.inherit = False; sh.name = name
        tf = sh.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.JUSTIFY
        runs = quote_runs if isinstance(quote_runs, list) else [(quote_runs, {})]
        r = p.add_run(); r.text = "“"; self._style_run(r, size=28, bold=True, color="ink", font=self.body_font)
        for t, o in runs:
            r = p.add_run(); r.text = t
            oo = dict(size=20, italic=True, color="ink", font=self.body_font); oo.update(o)
            self._style_run(r, **oo)
        r = p.add_run(); r.text = "”"; self._style_run(r, size=28, bold=True, color="ink", font=self.body_font)
        return sh

    def numbered_circle(self, slide, x, y, d, label, filled=False, color="signal-red", size=36, name=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, _in(x), _in(y), _in(d), _in(d))
        sh.shadow.inherit = False
        c = COLORS.get(color, color)
        if filled: sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(c)
        else: sh.fill.solid(); sh.fill.fore_color.rgb = _rgb("FFFFFF")
        sh.line.color.rgb = _rgb(c); sh.line.width = Pt(3)
        tf = sh.text_frame; tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(label)
        self._style_run(r, size=size, bold=True, color="FFFFFF" if filled else c, font=self.body_font)
        if name: sh.name = name
        return sh

    def agenda_slide(self, items: Sequence[str], current: int | None = None, number=None):
        """Agenda 'arc': a large crimson circle bleeding off the left edge, numbered circles on it and
        white boxes with a 3pt #C00000 frame and BOLD CAPS labels. `current` (1-based) is a filled
        circle named '!!Agenda current' so Morph slides it between agenda slides."""
        s = self.new_slide("blank")
        arc = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(-5029200), Emu(-751840), Emu(8514080), Emu(8361680))
        arc.fill.background(); arc.line.color.rgb = _rgb("C00000"); arc.line.width = Pt(6)
        arc.shadow.inherit = False; arc.name = "Agenda arc"
        n = len(items)
        # positions on the arc (from the source deck for 4 items, interpolated otherwise)
        top, bottom = 0.283, 6.136
        cx_r = 3.0  # arc center x (in) ~ -5029200+8514080/2 → -0.84in ; radius ≈ 4.6in
        import math
        cxa, cya, ra = (-5029200 + 8514080 / 2) / 914400, (-751840 + 8361680 / 2) / 914400, 8514080 / 2 / 914400
        circles = []
        for i, it in enumerate(items):
            yc = top + 0.59 + (bottom - top) * (i / max(n - 1, 1))
            dy = yc - cya
            xc = cxa + math.sqrt(max(ra * ra - dy * dy, 0))
            c = self.numbered_circle(s, xc - 0.59, yc - 0.59, 1.18, i + 1, filled=False, name=f"Agenda num {i+1}")
            circles.append((xc, yc))
            box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, _in(xc + 0.9), _in(yc - 0.52), _in(7.86), _in(1.04))
            box.fill.solid(); box.fill.fore_color.rgb = _rgb("FFFFFF")
            box.line.color.rgb = _rgb("C00000"); box.line.width = Pt(3); box.shadow.inherit = False
            box.name = f"Agenda item {i+1}"
            tf = box.text_frame; tf.margin_left = Inches(0.45)
            p = tf.paragraphs[0]; r = p.add_run(); r.text = it.upper()
            self._style_run(r, size=30 if n <= 4 else 24, bold=True, color="ink", font=self.body_font)
        if current:
            xc, yc = circles[current - 1]
            self.numbered_circle(s, xc - 0.59, yc - 0.59, 1.18, current, filled=True, name="!!Agenda current")
        if number is not None:
            self.slide_number(s, number)
        return s

    def stat_bubble(self, slide, x, y, d, value, label, fill="FFF2CC", line="F9B702", value_color="7E0029", name=None):
        """Pastel circle with a big number and a small label (dataset facts: '30 SaaS', '162 pricings')."""
        sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, _in(x), _in(y), _in(d), _in(d))
        sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(fill)
        sh.line.color.rgb = _rgb(line); sh.line.width = Pt(1.5); sh.shadow.inherit = False
        tf = sh.text_frame; tf.margin_left = tf.margin_right = 0; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        vsize = int(min(d * 20, d * 72 * 0.72 / (0.6 * max(len(str(value)), 1))))
        r = p.add_run(); r.text = str(value); self._style_run(r, size=max(vsize, 14), bold=True, color=value_color, font=self.body_font)
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r = p2.add_run(); r.text = label; self._style_run(r, size=max(int(d * 8), 11), color="ink-soft", font=self.body_font)
        if name: sh.name = name
        return sh

    def arrow(self, slide, x, y, w, h, color="crimson-deep", direction="right", text=None, name=None):
        geom = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW, "left": MSO_SHAPE.LEFT_ARROW, "up": MSO_SHAPE.UP_ARROW}[direction]
        sh = slide.shapes.add_shape(geom, _in(x), _in(y), _in(w), _in(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(COLORS.get(color, color))
        sh.line.fill.background(); sh.shadow.inherit = False
        if text:
            p = sh.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = text; self._style_run(r, size=12, bold=True, color="white", font=self.body_font)
        if name: sh.name = name
        return sh

    def rq_row(self, slide, x, y, w, tag, label, question, note="", name=None):
        """Research-question row: rounded white card, small crimson 'RQ1' pill, bold label, question."""
        h = 0.62
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _in(x), _in(y), _in(w), _in(h))
        card.adjustments[0] = 0.12
        card.fill.solid(); card.fill.fore_color.rgb = _rgb("FFFFFF")
        card.line.color.rgb = _rgb("F8B2CD"); card.line.width = Pt(1.25); card.shadow.inherit = False
        card.name = (name or f"RQ {tag}") + " card"
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _in(x + 0.15), _in(y + 0.17), _in(0.62), _in(0.28))
        pill.fill.solid(); pill.fill.fore_color.rgb = _rgb("7E0029"); pill.line.fill.background(); pill.shadow.inherit = False
        pill.name = (name or f"RQ {tag}") + " pill"
        tf = pill.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = tag; self._style_run(r, size=11, bold=True, color="white", font=self.body_font)
        self.text(slide, x + 0.92, y + 0.12, 1.8, 0.4, label, size=13, bold=True, color="crimson-deep", anchor="middle",
                  name=(name or f"RQ {tag}") + " label")
        self.text(slide, x + 2.8, y + 0.06, w - 2.95, 0.5,
                  [[(question, {"bold": True, "size": 13})]] + ([[(note, {"size": 10, "color": "7F7F7F"})]] if note else []),
                  size=13, color="ink", anchor="middle", name=(name or f"RQ {tag}") + " text")
        return card

    # ---------- dynamics --------------------------------------------------------
    def build_step(self, slide, morph=True, option="byObject", dur_ms=None):
        """Duplicate `slide` right after itself and give the copy a Morph transition — the group's
        main storytelling device (progressive builds, moving highlights, zoom-ins)."""
        new = duplicate_slide(self.prs, slide)
        if morph:
            set_transition(new, "morph", option=option, dur_ms=dur_ms)
        return new

    def morph(self, slide, option="byObject", dur_ms=None, speed="slow"):
        set_transition(slide, "morph", option=option, dur_ms=dur_ms, speed=speed)

    def fade(self, slide, speed="slow"):
        set_transition(slide, "fade", speed=speed)

    def animate(self, slide, shape, effect="wipe", direction="from_left", trigger="click", delay=0, dur=None, cls="entr"):
        """Append one effect to the slide's timeline (rebuilds the <p:timing>)."""
        lst = getattr(slide.part, "_isa_steps", None)
        if lst is None:
            lst = []
            slide.part._isa_steps = lst
        st = dict(shape=shape, effect=effect, direction=direction, trigger=trigger, delay=delay, cls=cls)
        if dur is not None: st["dur"] = dur
        lst.append(st)
        apply_animations(slide, [dict(x) for x in lst])

    def animate_sequence(self, slide, shapes, effect="zoom", first="click", then="after", **kw):
        for i, sh in enumerate(shapes):
            self.animate(slide, sh, effect, trigger=first if i == 0 else then, **kw)

    def hide(self, slide, shape):
        """Hide a shape (used on build slides: keep it for Morph matching but invisible)."""
        shape._element.find(".//p:cNvPr", NS).set("hidden", "1")

    def move_offslide(self, shape, side="top"):
        """Park a shape outside the canvas so Morph flies it in on the next slide."""
        if side == "top": shape.top = Emu(-shape.height - 200000)
        elif side == "bottom": shape.top = Emu(SLIDE_H + 200000)
        elif side == "left": shape.left = Emu(-shape.width - 200000)
        else: shape.left = Emu(SLIDE_W + 200000)

    # ---------- component library -------------------------------------------------
    library_dir = os.path.join(HERE, "library")

    def list_components(self, category=None):
        """Catalogue of reusable components (kit/library/catalog.json): id, category, name,
        description, number of slides, transitions per slide and effects used."""
        import json
        cat = json.load(open(os.path.join(self.library_dir, "catalog.json"), encoding="utf-8"))
        return [c for c in cat if category is None or c["category"] == category]

    def component_path(self, cid):
        p = os.path.join(self.library_dir, cid + ".pptx")
        if os.path.exists(p):
            return p
        if os.path.exists(p + ".b64.txt"):
            return restore_pptx(p + ".b64.txt", p)
        raise FileNotFoundError(f"component {cid!r}: download kit/library/{cid}.pptx.b64.txt next to isa_deck.py/library/")

    def component_texts(self, cid):
        """Every distinct text run in a component — the exact strings `replacements` can target
        (text shown in caps may be stored in lower case)."""
        from pptx import Presentation as _P
        seen = []
        for sl in _P(self.component_path(cid)).slides:
            for t in sl._element.iter("{%s}t" % NS["a"]):
                if t.text and t.text.strip() and t.text not in seen:
                    seen.append(t.text)
        return seen

    def insert_component(self, cid, replacements=None, number=None, slides=None):
        """Append a library component — the group's ORIGINAL slides with their media, Morph
        transitions and animation timelines — and return the new slides.
        replacements: {"old text": "new text"} applied to every run (e.g. {"BASIC": "FREE", "$10.99": "$0"}).
        number: logical slide number written in the bottom-right number box of every slide.
        slides: optional 1-based subset of the component's slides (e.g. [1, 2, 5])."""
        new = import_slides(self.prs, self.component_path(cid), set(slides) if slides else None)
        replace_text(new, replacements or {})
        if number is not None:
            set_logical_number(new, number)
        return new

    # ---------- output ----------------------------------------------------------
    def save(self, path):
        # drop the hidden seed slide holding the logo library
        seeds = [x for x in [self._seed] + list(self._seeds.values()) if x is not None]
        seed_ids = {x.slide_id for x in seeds}
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            if int(sid.get("id")) in seed_ids:
                lst.remove(sid)
                self.prs.part.drop_rel(sid.rId)
        self._seed, self._seeds = None, {}
        self.prs.save(path)
        return path


# ----------------------------------------------------------------------------
# Post-processing any .pptx (e.g. a deck exported from Claude Design / Slides)
# ----------------------------------------------------------------------------
def apply_spec(pptx_in, spec, pptx_out=None):
    """Add transitions + animations to an existing deck, addressing shapes by NAME.

    spec = {"slides": {"3": {"transition": {"kind": "morph", "option": "byObject", "dur_ms": 2000},
                              "animations": [{"shape": "Layer access", "effect": "wipe",
                                              "direction": "from_left", "trigger": "click"}]},
                       "4": {"duplicate_of": 3}}}          # (duplicate_of is applied first)
    Slide keys are 1-based positions in the ORIGINAL deck."""
    import json
    if isinstance(spec, str):
        spec = json.load(open(spec, encoding="utf-8"))
    prs = Presentation(pptx_in)
    slides = list(prs.slides)
    for key, cfg in spec.get("slides", {}).items():
        s = slides[int(key) - 1]
        tr = cfg.get("transition")
        if tr:
            set_transition(s, tr.get("kind", "morph"), option=tr.get("option", "byObject"), speed=tr.get("speed", "slow"),
                           dur_ms=tr.get("dur_ms"), advance_after_ms=tr.get("advance_after_ms"))
        anims = cfg.get("animations")
        if anims:
            steps = []
            for a in anims:
                a = dict(a)
                a["shape"] = find_shape(s, a["shape"])
                steps.append(a)
            apply_animations(s, steps)
    prs.save(pptx_out or pptx_in)
    return pptx_out or pptx_in


if __name__ == "__main__":
    import sys
    if len(sys.argv) >= 3 and sys.argv[1] == "--restore":
        print(restore_pptx(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None))
        sys.exit(0)
    if len(sys.argv) < 3:
        print("usage: python isa_deck.py deck.pptx spec.json [out.pptx]   (adds transitions/animations by shape name)")
        sys.exit(1)
    print(apply_spec(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None))


__all__ = ["apply_spec", "restore_pptx", "import_slide", "import_slides", "replace_text", "set_logical_number", "Deck", "apply_animations", "set_transition", "duplicate_slide", "find_shape", "COLORS", "CARD_STYLES",
           "ENTRANCE_EFFECTS", "EXIT_EFFECTS"]
