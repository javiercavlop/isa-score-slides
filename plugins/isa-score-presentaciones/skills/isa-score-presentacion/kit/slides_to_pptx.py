# ISA · SCORE Slides — (c) 2026 Francisco Javier Cavero López y Alejandro García Fernández
# (SCORE Lab, I3US, Universidad de Sevilla; Grupo ISA). Licencia CC BY 4.0 — ver LICENSE.
# Las presentaciones generadas con este kit pueden usarse libremente sin atribución (LICENSE, parte 2).
# ISA · SCORE Slides — slides_to_pptx.py
# Convierte un deck de Claude Slides (project/deck.json + project/slides/*.html) en un .pptx NATIVO
# del grupo: formas, textos e imágenes editables, Morph donde el deck usa "magic move" y animaciones
# de entrada donde usa data-build-in. Usa isa_deck.py (plantilla, portada/cierre reales y biblioteca).
#
#   pip install python-pptx lxml playwright && python -m playwright install chromium
#   npm i @fontsource/roboto @fontsource/roboto-condensed lucide-static   (fuentes e iconos para medir)
#   python slides_to_pptx.py DECK_DIR out.pptx --assets ASSETS_DIR [--map map.json] [--node node_modules]
#
# DECK_DIR: carpeta con project/deck.json y project/slides/ (lo que deja un `read` del deck).
# ASSETS_DIR: imágenes del deck; cada /_blob/<id> se resuelve como <id>.* o vía ASSETS_DIR/assets.json {id: fichero}.
# map.json (opcional): portada/cierre reales y componentes originales de la biblioteca en lugar de
# diapositivas concretas (ver pptx-workflow.md, "Camino C").
import argparse, copy, glob, json, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from isa_deck import Deck, set_transition, apply_animations, NS, P
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from lxml import etree

E = 6350  # EMU por px (lienzo 1920 px = 13,333 in)
ICON = {'Clock': 'clock', 'Lock': 'lock', 'Warning': 'triangle-alert', 'Star': 'star', 'Settings': 'settings',
        'Activity': 'activity', 'Wrench': 'wrench', 'Code': 'code', 'Verified': 'badge-check', 'Users': 'users',
        'CheckCircle': 'circle-check', 'Check': 'check', 'Globe': 'globe', 'Book': 'book-open', 'Tool': 'hammer',
        'Lightbulb': 'lightbulb', 'Link': 'link', 'Search': 'search', 'Chart': 'chart-column', 'Cloud': 'cloud',
        'Database': 'database', 'Key': 'key', 'Lightning': 'zap', 'PaperPlane': 'send', 'Play': 'play',
        'Home': 'house', 'GraduationCap': 'graduation-cap', 'ThumbsUp': 'thumbs-up', 'Trust': 'shield-check', 'Chat': 'message-circle'}


# ---------------------------------------------------------------------------- 1. Render + extract
def font_css(node):
    fs = os.path.join(node, '@fontsource')
    rows = [('Roboto', 300, 'normal', 'roboto/files/roboto-latin-300-normal.woff2'),
            ('Roboto', 400, 'normal', 'roboto/files/roboto-latin-400-normal.woff2'),
            ('Roboto', 400, 'italic', 'roboto/files/roboto-latin-400-italic.woff2'),
            ('Roboto', 700, 'normal', 'roboto/files/roboto-latin-700-normal.woff2'),
            ('Roboto', 700, 'italic', 'roboto/files/roboto-latin-700-italic.woff2'),
            ('Roboto Condensed', 400, 'normal', 'roboto-condensed/files/roboto-condensed-latin-400-normal.woff2'),
            ('Roboto Condensed', 700, 'normal', 'roboto-condensed/files/roboto-condensed-latin-700-normal.woff2')]
    return ''.join(f"@font-face{{font-family:'{f}';font-weight:{w};font-style:{st};src:url(file://{fs}/{p})}}"
                   for f, w, st, p in rows if os.path.exists(os.path.join(fs, p)))


def icon_svg(node, name, color, w, h):
    f = os.path.join(node, 'lucide-static', 'icons', ICON.get(name, 'circle') + '.svg')
    if not os.path.exists(f):
        s = int(min(w, h) * 0.62)
        return f'<svg width="{s}" height="{s}" viewBox="0 0 24 24"><circle cx="12" cy="12" r="9" fill="none" stroke="{color}" stroke-width="2"/></svg>'
    svg = re.sub(r'<!--.*?-->', '', open(f).read(), flags=re.S)
    s = int(min(w, h) * 0.62)
    svg = re.sub(r'width="24"', f'width="{s}"', svg, count=1)
    svg = re.sub(r'height="24"', f'height="{s}"', svg, count=1)
    return svg.replace('currentColor', color)


def resolve_asset(assets, blob):
    m = os.path.join(assets, 'assets.json')
    if os.path.exists(m):
        mp = json.load(open(m))
        if blob in mp:
            return os.path.join(assets, mp[blob])
    hits = glob.glob(os.path.join(assets, blob + '.*'))
    if hits:
        return hits[0]
    raise FileNotFoundError(f'asset {blob} not found in {assets}')


def prep(src, deck_dir, assets, node):
    src = re.sub(r'<aside>.*?</aside>', '', src, flags=re.S)
    src = re.sub(r'src="/?_blob/([0-9a-f]{32})"', lambda m: f'src="file://{resolve_asset(assets, m.group(1))}"', src)
    src = re.sub(r'src="(project/ds/[^"]+)"', lambda m: f'src="file://{os.path.join(deck_dir, m.group(1))}"', src)

    def xi(m):
        attrs = m.group(1)
        name = re.search(r'name="(\w+)"', attrs).group(1)
        st = re.search(r'style="([^"]*)"', attrs).group(1)
        col = (re.search(r'color:(#[0-9a-fA-F]{6})', st) or [0, '#141414'])[1]
        w = int((re.search(r'width:(\d+)px', st) or [0, 48])[1]); h = int((re.search(r'height:(\d+)px', st) or [0, 48])[1])
        return (f'<span data-xicon="1" style="display:inline-flex;align-items:center;justify-content:center;flex:none;'
                f'width:{w}px;height:{h}px">{icon_svg(node, name, col, w, h)}</span>')
    src = re.sub(r'<x-icon([^>]*)></x-icon>', xi, src)

    def xs(m):
        attrs = m.group(1)
        kind = re.search(r'kind="([\w-]+)"', attrs).group(1)
        st = re.search(r'style="([^"]*)"', attrs).group(1)
        rest = re.sub(r'kind="[\w-]+"|style="[^"]*"', '', attrs)
        return f'<div data-xshape="{kind}" style="{st};display:block;flex:none" {rest}></div>'
    src = re.sub(r'<x-shape([^>]*)></x-shape>', xs, src)
    src = re.sub(r'<x-connector([^>]*)></x-connector>', lambda m: f'<i data-xconn="1"{m.group(1)}></i>', src)
    return src


JS = r'''
() => {
const out = [];
const px = v => parseFloat(v) || 0;
const visible = c => c && c !== 'transparent' && !/rgba\(\s*0,\s*0,\s*0,\s*0\s*\)/.test(c);
let rk = 0;
for (const sec of document.querySelectorAll('section[data-slide]')) {
  const S = sec.getBoundingClientRect();
  const rel = r => ({x: r.left - S.left, y: r.top - S.top, w: r.width, h: r.height});
  const slide = {id: sec.dataset.slide, transition: sec.getAttribute('data-transition'), items: []};
  const textBlock = el => ['P','H1','H2','H3','LI'].includes(el.tagName);
  function runsOf(el) {
    // per character positions to keep the browser's line breaks
    const runs = [];
    const walker = document.createTreeWalker(el, NodeFilter.SHOW_TEXT | NodeFilter.SHOW_ELEMENT);
    let node; let lastTop = null;
    const range = document.createRange();
    while ((node = walker.nextNode())) {
      if (node.nodeType === 1) { if (node.tagName === 'BR') { runs.push({br: true}); lastTop = null; } continue; }
      const p = node.parentElement; const cs = getComputedStyle(p);
      const style = {color: cs.color, weight: cs.fontWeight, italic: cs.fontStyle === 'italic',
                     underline: cs.textDecorationLine.includes('underline'), strike: cs.textDecorationLine.includes('line-through'),
                     size: px(cs.fontSize), family: cs.fontFamily, transform: cs.textTransform, ls: cs.letterSpacing};
      const t = node.data;
      let buf = '';
      for (let i = 0; i < t.length; i++) {
        range.setStart(node, i); range.setEnd(node, i + 1);
        const rs = range.getClientRects();
        const top = rs.length ? Math.round(rs[0].top) : lastTop;
        if (lastTop !== null && top !== null && Math.abs(top - lastTop) > style.size * 0.5) {
          if (buf) runs.push({text: buf, ...style}); buf = '';
          runs.push({br: true, soft: true});
          if (t[i] === ' ') { lastTop = top; continue; }
        }
        if (top !== null) lastTop = top;
        buf += t[i];
      }
      if (buf) runs.push({text: buf, ...style});
    }
    // collapse whitespace like the browser
    for (const r of runs) if (r.text !== undefined) r.text = r.text.replace(/\s+/g, ' ');
    // trim spaces around soft breaks
    for (let i = 0; i < runs.length; i++) if (runs[i].br) {
      if (i > 0 && runs[i-1].text) runs[i-1].text = runs[i-1].text.replace(/ $/, '');
      if (i + 1 < runs.length && runs[i+1].text) runs[i+1].text = runs[i+1].text.replace(/^ /, '');
    }
    return runs.filter(r => r.br || r.text !== '');
  }
  function textExtent(el) {
    const range = document.createRange(); range.selectNodeContents(el);
    const r = range.getBoundingClientRect(); return rel(r);
  }
  function boxOf(el, cs, prims) {
    const r = rel(el.getBoundingClientRect());
    const bw = ['Top','Right','Bottom','Left'].map(s => px(cs['border' + s + 'Width']));
    const bc = ['Top','Right','Bottom','Left'].map(s => cs['border' + s + 'Color']);
    const bs = ['Top','Right','Bottom','Left'].map(s => cs['border' + s + 'Style']);
    const bg = cs.backgroundColor; const sh = cs.boxShadow;
    const hasB = bw.some((w, i) => w > 0 && bs[i] !== 'none' && visible(bc[i]));
    if (!(visible(bg) || hasB || (sh && sh !== 'none'))) return;
    prims.push({type: 'box', ...r, bg: visible(bg) ? bg : null, bw, bc, bs, radius: cs.borderTopLeftRadius,
                shadow: sh !== 'none' ? sh : null, opacity: px(cs.opacity)});
  }
  function walk(el, prims) {
    if (el.tagName === 'svg') { const id = 'rk' + (rk++); el.setAttribute('data-rk', id); prims.push({type: 'raster', key: id, ...rel(el.getBoundingClientRect())}); return; }
    if (el.dataset && el.dataset.xicon) { const sv = el.querySelector('svg'); const id = 'rk' + (rk++); sv.setAttribute('data-rk', id); prims.push({type: 'raster', key: id, ...rel(sv.getBoundingClientRect())}); return; }
    if (el.tagName === 'IMG') {
      const r = rel(el.getBoundingClientRect()); const cs = getComputedStyle(el);
      boxOf(el, cs, prims);
      let x = r.x, y = r.y, w = r.w, h = r.h;
      const bwid = px(cs.borderLeftWidth);
      x += bwid; y += bwid; w -= 2*bwid; h -= 2*bwid;
      if (el.naturalWidth && cs.objectFit === 'contain') {
        const s = Math.min(w / el.naturalWidth, h / el.naturalHeight);
        const nw = el.naturalWidth * s, nh = el.naturalHeight * s; x += (w - nw) / 2; y += (h - nh) / 2; w = nw; h = nh;
      }
      prims.push({type: 'img', src: el.getAttribute('src'), x, y, w, h, opacity: px(cs.opacity), alt: el.alt});
      return;
    }
    if (el.dataset && el.dataset.xshape) { const cs = getComputedStyle(el); prims.push({type: 'shape', kind: el.dataset.xshape, ...rel(el.getBoundingClientRect()), bg: cs.backgroundColor}); return; }
    if (el.tagName === 'I') return;
    const cs = getComputedStyle(el);
    if (cs.display === 'none') return;
    boxOf(el, cs, prims);
    if (textBlock(el)) {
      const ext = textExtent(el);
      const r = rel(el.getBoundingClientRect());
      const pl = px(cs.paddingLeft) + px(cs.borderLeftWidth), pr = px(cs.paddingRight) + px(cs.borderRightWidth);
      prims.push({type: 'text', x: r.x + pl, w: r.w - pl - pr, y: ext.y, h: ext.h, align: cs.textAlign,
                  lh: cs.lineHeight === 'normal' ? px(cs.fontSize) * 1.2 : px(cs.lineHeight), size: px(cs.fontSize),
                  runs: runsOf(el), tag: el.tagName});
      return;
    }
    for (const c of el.children) walk(c, prims);
  }
  let n = 0;
  for (const el of sec.children) {
    n++;
    if (el.tagName === 'I' && el.dataset.xconn) {
      const g = a => el.getAttribute(a);
      const cs = el.getAttribute('style') || '';
      const col = (cs.match(/color:(#[0-9a-fA-F]{6})/) || [0, '#7e0029'])[1];
      const bw = (cs.match(/border-width:(\d+)px/) || [0, '2'])[1];
      slide.items.push({idx: n, id: el.id || null, build: g('data-build-in'), prims: [{type: 'line', x1: +g('x1'), y1: +g('y1'), x2: +g('x2'), y2: +g('y2'),
        head: g('head') || 'end', route: g('route') || 'straight', color: col, width: +bw, dashed: /dashed/.test(cs)}]});
      continue;
    }
    const prims = []; walk(el, prims);
    slide.items.push({idx: n, id: el.id || null, build: el.getAttribute('data-build-in'), tag: el.tagName, prims});
  }
  out.push(slide);
}
return out;
}
'''


def extract(deck_dir, ids, assets, node, work):
    from playwright.sync_api import sync_playwright
    sl = os.path.join(deck_dir, 'project', 'slides')
    parts = []
    for sid in ids:
        src = prep(open(os.path.join(sl, sid + '.html')).read(), deck_dir, assets, node)
        parts.append(src.replace('<section ', f'<section data-slide="{sid}" ', 1))
    page = ('<!doctype html><html><head><meta charset="utf-8"><style>' + font_css(node) +
            '*{margin:0;padding:0;box-sizing:border-box} body{background:#ddd}'
            'section{position:relative;width:1920px;height:1080px;overflow:hidden;margin-bottom:20px}'
            'section > *{position:absolute} h1,h2,h3{font-weight:600} ul,ol{padding-left:1.2em} i[data-xconn]{display:none}'
            '</style></head><body>' + ''.join(parts) + '</body></html>')
    os.makedirs(os.path.join(work, 'raster'), exist_ok=True)
    open(os.path.join(work, 'render.html'), 'w').write(page)
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={'width': 1920, 'height': 1080})
        pg.goto('file://' + os.path.join(work, 'render.html'))
        pg.wait_for_timeout(1500)
        pg.evaluate('document.fonts.ready')
        data = pg.evaluate(JS)
        keys = pg.evaluate("() => [...document.querySelectorAll('[data-rk]')].map(e => [e.getAttribute('data-rk'), e.outerHTML, e.getBoundingClientRect().width, e.getBoundingClientRect().height])")
        rp = b.new_page(device_scale_factor=2)
        for k, outer, w, h in keys:
            rp.set_viewport_size({'width': max(1, int(w + 1)), 'height': max(1, int(h + 1))})
            rp.set_content(f'<html><body style="margin:0;background:transparent">{outer}</body></html>')
            rp.query_selector('svg').screenshot(path=os.path.join(work, 'raster', k + '.png'), omit_background=True)
        b.close()
    return {s['id']: s for s in data}


# ---------------------------------------------------------------------------- 2. Draw native shapes
def emu(v): return Emu(int(round(v * E)))

def parse_color(c):
    if not c: return None, 1.0
    if c.startswith('#'):
        return c[1:7].upper(), 1.0
    m = re.match(r'rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)', c)
    if not m: return None, 1.0
    r, g, b = (int(float(m.group(i))) for i in (1, 2, 3))
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return f'{r:02X}{g:02X}{b:02X}', a

def set_alpha(fill_or_line_color_el, a):
    if a >= 0.999: return
    srgb = fill_or_line_color_el
    al = etree.SubElement(srgb, '{%s}alpha' % NS['a']); al.set('val', str(int(a * 100000)))

def solid_fill(shape, color):
    hexc, a = parse_color(color)
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(hexc)
    if a < 1:
        set_alpha(shape.fill._xPr.find('.//a:srgbClr', NS), a)

def no_line(shape):
    shape.line.fill.background()

def set_line(shape, color, width_px, dashed=False):
    hexc, a = parse_color(color)
    shape.line.color.rgb = RGBColor.from_string(hexc)
    shape.line.width = emu(width_px)
    if dashed: shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if a < 1:
        set_alpha(shape.line._ln.find('.//a:srgbClr', NS), a)

def radius_px(r, w, h):
    if r is None: return 0
    if r.endswith('%'): return min(w, h) * float(r[:-1]) / 100
    return float(r.split()[0].replace('px', '') or 0)

def add_geom(shapes, x, y, w, h, rad):
    if rad >= min(w, h) / 2 - 0.5 and rad > 0:
        if abs(w - h) < 1:
            return shapes.add_shape(MSO_SHAPE.OVAL, emu(x), emu(y), emu(w), emu(h))
        sh = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h)); sh.adjustments[0] = 0.5; return sh
    if rad > 0:
        sh = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
        sh.adjustments[0] = min(0.5, rad / min(w, h)); return sh
    return shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))

def clean(sh):
    # remove default text style / shadow from theme for autoshapes
    sh.shadow.inherit = False
    if sh.has_text_frame:
        sh.text_frame.text = ''

def draw_box(shapes, p):
    x, y, w, h = p['x'], p['y'], p['w'], p['h']
    if w < 0.5 or h < 0.5: return
    rad = radius_px(p['radius'], w, h)
    # spread ring shadows (0 0 0 Npx color)
    if p.get('shadow'):
        for part in re.findall(r'(rgba?\([^)]*\))\s+(-?[\d.]+)px\s+(-?[\d.]+)px\s+([\d.]+)px\s+([\d.]+)px', p['shadow']):
            col, dx, dy, blur, spread = part[0], float(part[1]), float(part[2]), float(part[3]), float(part[4])
            if blur == 0 and spread > 0:
                s = spread
                ring = add_geom(shapes, x - s, y - s, w + 2 * s, h + 2 * s, rad + s if rad > 0 else 0)
                clean(ring); solid_fill(ring, col); no_line(ring)
    bw, bc, bs = p['bw'], p['bc'], p['bs']
    vis = [bw[i] > 0 and bs[i] != 'none' and parse_color(bc[i])[0] is not None and parse_color(bc[i])[1] > 0 for i in range(4)]
    uniform = all(vis) and len(set(bw)) == 1 and len(set(bc)) == 1
    if p['bg'] or uniform:
        if uniform:
            b = bw[0]
            sh = add_geom(shapes, x + b / 2, y + b / 2, w - b, h - b, max(0, rad - b / 2))
            clean(sh)
            if p['bg']: solid_fill(sh, p['bg'])
            else: sh.fill.background()
            set_line(sh, bc[0], b, bs[0] == 'dashed')
        else:
            sh = add_geom(shapes, x, y, w, h, rad); clean(sh); solid_fill(sh, p['bg']); no_line(sh)
    if not uniform:
        for i, side in enumerate(('t', 'r', 'b', 'l')):
            if not vis[i]: continue
            b = bw[i]
            rx, ry, rw, rh = {'t': (x, y, w, b), 'b': (x, y + h - b, w, b), 'l': (x, y, b, h), 'r': (x + w - b, y, b, h)}[side]
            bar = shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(rx), emu(ry), emu(rw), emu(rh)); clean(bar)
            solid_fill(bar, bc[i]); no_line(bar)
            if bs[i] == 'dashed':
                bar.fill.background(); set_line(bar, bc[i], b, True)

def font_for(run):
    fam = run.get('family', '')
    w = int(float(run.get('weight', 400)))
    if 'Condensed' in fam: return 'Arial Narrow', w >= 600
    if w <= 300: return 'Roboto Light', False
    return 'Roboto', w >= 600

def draw_text(shapes, p):
    runs = p['runs']
    if not runs: return
    size = p['size']; lh = p['lh']
    lines = 1 + sum(1 for r in runs if r.get('br'))
    content_line = size * 1.17
    pad = max(0.0, (lh - content_line) / 2)
    y = p['y'] - pad
    h = max(p['h'] + 2 * pad, lines * lh)
    x, w = p['x'], p['w']
    extra = max(12, w * 0.06)
    al = p['align']
    if al in ('center', '-webkit-center'):
        x -= extra / 2
    elif al in ('right', 'end'):
        x -= extra
    w += extra
    tb = shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para = tf.paragraphs[0]
    para.alignment = {'center': PP_ALIGN.CENTER, '-webkit-center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT,
                      'justify': PP_ALIGN.JUSTIFY}.get(al, PP_ALIGN.LEFT)
    para.line_spacing = Pt(lh * 0.5)
    for r in runs:
        if r.get('br'):
            para.add_line_break(); continue
        t = r['text']
        if r.get('transform') == 'uppercase': t = t.upper()
        run = para.add_run(); run.text = t
        f = run.font
        name, bold = font_for(r)
        f.name = name; f.bold = bold; f.size = Pt(r['size'] * 0.5)
        f.italic = bool(r.get('italic'))
        if r.get('underline'): f.underline = True
        hexc, a = parse_color(r['color'])
        if hexc: f.color.rgb = RGBColor.from_string(hexc)
        rPr = run._r.get_or_add_rPr()
        if r.get('strike'): rPr.set('strike', 'sngStrike')
        ls = r.get('ls', 'normal')
        if ls and ls != 'normal':
            try: rPr.set('spc', str(int(float(ls.replace('px', '')) * 0.5 * 100)))
            except ValueError: pass
        # East-asian / symbol fallbacks keep the same font
        for tag in ('latin', 'cs'):
            el = rPr.find('{%s}%s' % (NS['a'], tag))
        ea = etree.SubElement(rPr, '{%s}ea' % NS['a']); ea.set('typeface', name)
        sym = etree.SubElement(rPr, '{%s}sym' % NS['a']); sym.set('typeface', name)
    return tb

def draw_line(shapes, p):
    x1, y1, x2, y2 = p['x1'], p['y1'], p['x2'], p['y2']
    if p['route'] == 'hv': pts = [(x1, y1), (x2, y1), (x2, y2)]
    elif p['route'] == 'vh': pts = [(x1, y1), (x1, y2), (x2, y2)]
    else: pts = [(x1, y1), (x2, y2)]
    if len(pts) == 2:
        sh = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    else:
        fb = shapes.build_freeform(emu(pts[0][0]), emu(pts[0][1]), scale=1.0)
        fb.add_line_segments([(emu(a), emu(b)) for a, b in pts[1:]], close=False)
        sh = fb.convert_to_shape(); sh.fill.background(); sh.shadow.inherit = False
    set_line(sh, p['color'], p['width'], p['dashed'])
    ln = sh.line._get_or_add_ln()
    if p['head'] in ('end', 'both'):
        te = etree.SubElement(ln, '{%s}tailEnd' % NS['a']); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    if p['head'] == 'both':
        he = etree.SubElement(ln, '{%s}headEnd' % NS['a']); he.set('type', 'triangle'); he.set('w', 'med'); he.set('len', 'med')
    # schema order: headEnd must precede tailEnd
    if p['head'] == 'both':
        te = ln.find('{%s}tailEnd' % NS['a']); ln.remove(te); ln.append(te)
    return sh

def img_path(src):
    return src.replace('file://', '')

def draw_prim(shapes, p):
    t = p['type']
    if t == 'box': draw_box(shapes, p)
    elif t == 'text': draw_text(shapes, p)
    elif t == 'img':
        pic = shapes.add_picture(img_path(p['src']), emu(p['x']), emu(p['y']), emu(p['w']), emu(p['h']))
        if p.get('alt'): pic._element.find('.//p:cNvPr', NS).set('descr', p['alt'])
    elif t == 'raster':
        shapes.add_picture(os.path.join(WORK, "raster", p["key"] + ".png"), emu(p['x']), emu(p['y']), emu(p['w']), emu(p['h']))
    elif t == 'shape':
        kind = {'arrow-right': MSO_SHAPE.RIGHT_ARROW, 'arrow-left': MSO_SHAPE.LEFT_ARROW, 'arrow-down': MSO_SHAPE.DOWN_ARROW,
                'arrow-up': MSO_SHAPE.UP_ARROW, 'ellipse': MSO_SHAPE.OVAL, 'diamond': MSO_SHAPE.DIAMOND}.get(p['kind'], MSO_SHAPE.RECTANGLE)
        sh = shapes.add_shape(kind, emu(p['x']), emu(p['y']), emu(p['w']), emu(p['h'])); clean(sh)
        solid_fill(sh, p['bg']); no_line(sh)
    elif t == 'line': draw_line(shapes, p)

EFFECT = {'fade': ('fade', None, 500), 'rise': ('float_in', None, 600), 'drop': ('wipe', 'from_top', 500),
          'left': ('wipe', 'from_left', 500), 'right': ('wipe', 'from_right', 500), 'scale': ('zoom', None, 500),
          'pop': ('zoom', None, 500)}


def notes_of(deck_dir, sid):
    m = re.search(r'<aside>(.*?)</aside>', open(os.path.join(deck_dir, 'project', 'slides', sid + '.html')).read(), re.S)
    return re.sub(r'<[^>]+>', '', m.group(1)).replace('&amp;', '&') if m else ''


def convert(d, data, sid, deck_dir, transition_in):
    """One Slides <section> -> one native slide. Each pinned child becomes ONE shape (or one group),
    named '!!<id>' when it has an id, so Morph pairs it with the same id on the next slide."""
    s = d.new_slide('blank')
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    steps = []
    for it in data['items']:
        prims = it['prims']
        if not prims:
            continue
        before = len(s.shapes)
        if len(prims) == 1:
            draw_prim(s.shapes, prims[0])
            made = list(s.shapes)[before:]
            if not made:
                continue
            shp = made[0] if len(made) == 1 else s.shapes.add_group_shape(made)
        else:
            grp = s.shapes.add_group_shape()
            for pr in prims:
                draw_prim(grp.shapes, pr)
            if len(grp.shapes) == 0:
                grp._element.getparent().remove(grp._element); continue
            shp = grp
        shp.name = ('!!' + it['id']) if it.get('id') else f"{(it.get('tag') or 'el').lower()} {it['idx']}"
        if it.get('build'):
            m = re.match(r'(\w+) order (\d+)( auto)?', it['build'])
            if not m:
                continue
            eff, direction, dur = EFFECT[m.group(1)]
            if prims[0]['type'] == 'line':
                eff, dur = 'wipe', 500
                direction = {'left': 'from_left', 'right': 'from_right', 'drop': 'from_top', 'rise': 'from_bottom'}.get(m.group(1), 'from_left')
            steps.append(dict(order=int(m.group(2)), auto=bool(m.group(3)), idx=it['idx'], shape=shp, effect=eff,
                              direction=direction or 'from_left', dur=dur))
    # Slides build steps -> PowerPoint timeline: a new order = new step (click, or After Previous when
    # 'auto'); several elements with the same order = With Previous.
    steps.sort(key=lambda z: (z['order'], z['idx']))
    seq, prev = [], None
    for st in steps:
        trig = 'with' if prev == st['order'] else ('after' if st['auto'] else 'click')
        prev = st['order']
        seq.append(dict(shape=st['shape'], effect=st['effect'], direction=st['direction'], trigger=trig, dur=st['dur']))
    if seq:
        apply_animations(s, seq)
    # Slides puts the transition on the slide that LEAVES; PowerPoint on the slide that ARRIVES.
    if transition_in == 'magic':
        set_transition(s, 'morph')
    elif transition_in == 'fade':
        set_transition(s, 'fade')
    s.notes_slide.notes_text_frame.text = notes_of(deck_dir, sid)
    return s


def main():
    global WORK
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('deck_dir'); ap.add_argument('out')
    ap.add_argument('--assets', required=True)
    ap.add_argument('--map', default=None)
    ap.add_argument('--node', default=os.path.join(os.getcwd(), 'node_modules'))
    ap.add_argument('--work', default=os.path.join(os.getcwd(), '_s2p_work'))
    a = ap.parse_args()
    WORK = a.work
    os.makedirs(WORK, exist_ok=True)
    order = json.load(open(os.path.join(a.deck_dir, 'project', 'deck.json')))['order']
    mp = json.load(open(a.map)) if a.map else {}
    comps = {c['slide']: c for c in mp.get('components', [])}
    covered = {x for c in mp.get('components', []) for x in c.get('covers', [])}
    special = {mp.get('title_slide', {}).get('slide'), mp.get('closing_slide', {}).get('slide')} - {None}
    ids = [s for s in order if s not in comps and s not in covered and s not in special]
    data = extract(a.deck_dir, ids, a.assets, a.node, WORK)

    d = Deck()
    timings = {}
    prev_tr = None
    for sid in order:
        if sid in special and sid == mp.get('title_slide', {}).get('slide'):
            t = mp['title_slide']
            d.title_slide(t['title'], authors=[tuple(x) if isinstance(x, list) else x for x in t.get('authors', [])],
                          affiliation=t.get('affiliation', 'SCORE Lab, I3US, Universidad de Sevilla, Seville, Spain'),
                          venue=t.get('venue', ''), date=t.get('date', ''))
            prev_tr = None; continue
        if sid in special:
            t = mp['closing_slide']
            s = d.closing_slide(t['title'], authors=[tuple(x) if isinstance(x, list) else x for x in t.get('authors', [])],
                                email=t.get('email', ''), venue=t.get('venue', ''), date=t.get('date', ''),
                                qr_image=t.get('qr_image'), funding=t.get('funding', True))
            if not t.get('qr_image'):
                for sh in list(s.shapes):
                    if sh.name == 'QR placeholder':
                        sh._element.getparent().remove(sh._element)
            prev_tr = None; continue
        if sid in covered:
            continue
        if sid in comps:
            c = comps[sid]
            new = d.insert_component(c['id'], replacements=c.get('replacements'), number=c.get('number'), slides=c.get('slides'))
            for sl, nsid in zip(new, c.get('notes', [sid])):
                sl.notes_slide.notes_text_frame.text = notes_of(a.deck_dir, nsid)
            tm = new[0]._element.find(P + 'timing')
            if tm is not None and not c.get('timing_from'):
                timings[sid] = copy.deepcopy(tm)
            if c.get('timing_from') in timings:  # same component, same shape ids: reuse its animation tree
                el = new[0]._element
                for t in el.findall(P + 'timing'):
                    el.remove(t)
                ext = el.find(P + 'extLst')
                (ext.addprevious if ext is not None else el.append)(copy.deepcopy(timings[c['timing_from']]))
            prev_tr = None; continue
        convert(d, data[sid], sid, a.deck_dir, prev_tr or 'fade')
        prev_tr = data[sid].get('transition')
    d.save(a.out)
    print('saved', a.out, len(d.prs.slides), 'slides')


if __name__ == '__main__':
    main()
